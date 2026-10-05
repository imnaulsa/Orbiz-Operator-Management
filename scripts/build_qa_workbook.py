import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

def build_complete_qa_workbook(output_path="QA_Testing_Livestreaming_Production_V1.xlsx"):
    wb = openpyxl.Workbook()
    
    ws_summary = wb.active
    ws_summary.title = "QA Summary"
    ws_uat = wb.create_sheet(title="UAT Testing")
    ws_mobile = wb.create_sheet(title="UAT Mobile")
    
    FONT_FAMILY = "Rubik"
    
    # Styles
    f_title = Font(name=FONT_FAMILY, size=15, bold=True, color="1E293B")
    f_subtitle = Font(name=FONT_FAMILY, size=9.5, italic=True, color="475569")
    f_sec_header = Font(name=FONT_FAMILY, size=11, bold=True, color="1E293B")
    f_tbl_header = Font(name=FONT_FAMILY, size=9.5, bold=True, color="2D1F41")
    f_cell = Font(name=FONT_FAMILY, size=9, bold=False, color="1E293B")
    f_cell_bold = Font(name=FONT_FAMILY, size=9, bold=True, color="1E293B")
    f_cell_mono = Font(name="Consolas", size=9, bold=True, color="0F172A")
    f_kpi_label = Font(name=FONT_FAMILY, size=8.5, bold=True, color="475569")
    f_kpi_val = Font(name=FONT_FAMILY, size=13, bold=True, color="1E293B")
    
    fill_header = PatternFill(start_color="E8E2F0", end_color="E8E2F0", fill_type="solid")
    fill_header_sub = PatternFill(start_color="F1EBF9", end_color="F1EBF9", fill_type="solid")
    fill_kpi = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_zebra = PatternFill(start_color="FBFBFE", end_color="FBFBFE", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_bar = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_highlight = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
    
    fill_pass = PatternFill(start_color="DCF5E4", end_color="DCF5E4", fill_type="solid")
    font_pass = Font(name=FONT_FAMILY, size=9.5, bold=True, color="186938")
    fill_fail = PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid")
    font_fail = Font(name=FONT_FAMILY, size=9.5, bold=True, color="B02323")
    fill_blocked = PatternFill(start_color="FFEFCC", end_color="FFEFCC", fill_type="solid")
    font_blocked = Font(name=FONT_FAMILY, size=9.5, bold=True, color="915A00")
    fill_retest = PatternFill(start_color="DEEBFF", end_color="DEEBFF", fill_type="solid")
    font_retest = Font(name=FONT_FAMILY, size=9.5, bold=True, color="26569C")
    fill_not_tested = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_not_tested = Font(name=FONT_FAMILY, size=9.5, bold=False, color="475569")
    
    thin_border = Side(style='thin', color='CBD5E1')
    medium_border = Side(style='medium', color='94A3B8')
    double_bottom = Side(style='double', color='64748B')
    
    border_all = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    border_tbl_header = Border(left=thin_border, right=thin_border, top=medium_border, bottom=medium_border)
    border_total = Border(top=thin_border, bottom=double_bottom, left=thin_border, right=thin_border)
    
    al_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    al_center_top = Alignment(horizontal='center', vertical='top', wrap_text=True)
    al_left_top = Alignment(horizontal='left', vertical='top', wrap_text=True)
    al_right_top = Alignment(horizontal='right', vertical='top', wrap_text=True)
    al_left_center = Alignment(horizontal='left', vertical='center', wrap_text=True)
    al_right_center = Alignment(horizontal='right', vertical='center', wrap_text=True)

    # -------------------------------------------------------------
    # DATA DEFINITIONS: UAT Testing Test Cases
    # -------------------------------------------------------------
    uat_cases = [
        # Domain 1: Authentication & Session (AUTH)
        ("AUTH-01", "Authentication", "P0", "Super Admin", "All",
         "Login Super Admin ke Production Ops",
         "1. Buka URL login aplikasi.\n2. Masukkan email & password Super Admin aktif.\n3. Klik tombol Masuk.\n4. Buka navigasi ke /production.",
         "Super Admin berhasil login, diarahkan ke dashboard Production Ops, nama & role Super Admin termuat, serta memiliki akses ke semua menu dan selector lokasi.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-02", "Authentication", "P0", "Operator Manager", "Jakarta / Bandung",
         "Login Operator Manager & verifikasi menu",
         "1. Login dengan kredensial Operator Manager aktif (e.g. Jakarta).\n2. Amati landing page dan menu sidebar.",
         "Login sukses, masuk ke /production/overview, menu Ringkasan, Quotation (lihat), Master Data (lihat), Plotting Jadwal, Jadwal Livestreaming, Studio Timeline, dan Ringkasan Jadwal tersedia.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-03", "Authentication", "P0", "Operator Staff", "Jakarta / Bandung",
         "Login Operator Staff & verifikasi menu",
         "1. Login dengan kredensial Operator Staff aktif.\n2. Amati menu yang muncul di sidebar.",
         "Login sukses, diarahkan ke /production, menu yang tersedia terbatas pada Jadwal Livestreaming (read-only) dan Logbook Production Operator. Menu mutasi disembunyikan.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-04", "Authentication", "P0", "Host Manager", "Jakarta / Bandung",
         "Login Host Manager & verifikasi menu",
         "1. Login dengan kredensial Host Manager aktif.\n2. Amati navigasi sidebar.",
         "Login sukses, menu yang tersedia: Ringkasan, Plotting Jadwal, Jadwal Livestreaming, Ringkasan Jadwal, Studio Timeline, Availability Host, Perizinan Host, Jam & Logbook, dan Host Cost.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-05", "Authentication", "P0", "Host", "Jakarta / Bandung",
         "Login Host & direct redirect ke Jadwal Livestreaming",
         "1. Login dengan akun Host aktif.\n2. Amati routing dan landing page.",
         "Login sukses, otomatis diarahkan ke /production/schedule_table. Menu yang tersedia: Jadwal Livestreaming, Availability Host, Perizinan Host, dan Jam & Logbook.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-06", "Authentication", "P0", "Admin Sales", "All",
         "Login Admin Sales & direct redirect ke Quotation & Kuota",
         "1. Login dengan akun Admin Sales aktif.\n2. Amati routing dan landing page.",
         "Login sukses, otomatis diarahkan ke /production/quotation. Menu yang tersedia: Ringkasan dan Quotation & Kuota dengan hak akses penuh create/edit/delete quotation.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-07", "Authentication", "P0", "Semua Role", "All",
         "Proteksi route /production tanpa autentikasi",
         "1. Buka browser baru (Incognito/tanpa session).\n2. Langsung ketik URL /production atau /production/quotation.",
         "Aplikasi menolak akses dan me-redirect browser kembali ke halaman /login.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-08", "Authentication", "P0", "Semua Role", "All",
         "Login dengan password salah atau email tidak terdaftar",
         "1. Masukkan email valid dengan password salah.\n2. Coba juga email yang tidak terdaftar.",
         "Login ditolak dengan pesan error 'Invalid login credentials' tanpa membocorkan detail internal akun.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-09", "Authentication", "P0", "Semua Role", "All",
         "Session persistence saat refresh browser",
         "1. Login sebagai user aktif dan navigasi ke sub-halaman (e.g. /production/schedule_table).\n2. Tekan tombol Refresh browser (F5).",
         "Sesi tetap aktif, profil user tetap termuat, dan halaman tetap berada pada sub-halaman yang sedang dibuka tanpa kembali ke login.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-10", "Authentication", "P1", "Semua Role", "All",
         "Logout / Keluar mengakhiri session & history",
         "1. Klik tombol 'Keluar' pada sidebar.\n2. Setelah berada di /login, tekan tombol Back pada browser.",
         "Session auth berakhir, user tetap di /login, dan tombol Back tidak dapat membuka kembali halaman workspace /production.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-11", "Authentication", "P0", "Akun Nonaktif", "Jakarta / Bandung",
         "Login ditolak untuk akun nonaktif (active=false)",
         "1. Nonaktifkan akun uji di profiles (active=false).\n2. Coba login dengan email & password yang benar.",
         "Aplikasi menolak akses dan tidak mengizinkan masuk ke workspace production.",
         "Not Tested", "", "", "", ""),
        
        ("AUTH-12", "Authentication", "P0", "Akun Tanpa Profile", "All",
         "Login ditolak untuk user Auth tanpa record profiles",
         "1. Buat user di Auth tanpa profile terkait di database.\n2. Coba login.",
         "Akses ditolak secara aman, tidak ada data yang bocor, dan diarahkan keluar.",
         "Not Tested", "", "", "", ""),

        # Domain 2: Authorization & RBAC (RBAC)
        ("RBAC-01", "RBAC", "P0", "Host", "Jakarta",
         "Host dilarang akses Quotation & Master Data",
         "1. Login sebagai Host.\n2. Periksa sidebar navigasi.\n3. Coba akses via console / RPC langsung.",
         "Menu Quotation dan Master Data tidak muncul. Pemanggilan RPC production_manage ditolak dengan error 'Akses ditolak' (42501).",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-02", "RBAC", "P0", "Host", "Jakarta",
         "Host dilarang melakukan plotting jadwal & approval",
         "1. Login sebagai Host.\n2. Coba panggil RPC session_multi atau review_availability.",
         "RPC gagal dieksekusi dengan exception permission error (42501).",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-03", "RBAC", "P0", "Admin Sales", "All",
         "Admin Sales dilarang membuat Studio & Plotting",
         "1. Login sebagai Admin Sales.\n2. Coba panggil action 'studio' atau 'session_multi'.",
         "Aksi ditolak server-side dengan pesan error 'Hanya pengelola production' / 'Hanya Manager'.",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-04", "RBAC", "P0", "Admin Sales", "All",
         "Admin Sales tidak dapat melihat snapshot fee host",
         "1. Login sebagai Admin Sales.\n2. Buka data sesi live pada rentang aktif.",
         "Field 'host_fee' di-strip dari payload snapshot dan tidak ditampilkan di antarmuka.",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-05", "RBAC", "P0", "Operator Manager", "Jakarta",
         "Operator Manager akses read-only reference Quotation & Master Data",
         "1. Login sebagai Operator Manager.\n2. Buka menu Quotation & Kuota dan Master Data.",
         "Tabel referensi tampil (termasuk rate quotation dan studio kotanya), tetapi form tambah, tombol Edit, tombol Delete, dan tombol aksi mutasi tidak ditampilkan.",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-06", "RBAC", "P0", "Operator Manager", "Jakarta",
         "Operator Manager dilarang mutasi Master Brand & Quotation",
         "1. Panggil RPC production_manage dengan action 'brand_edit' atau 'quotation_delete' menggunakan akun Operator Manager.",
         "RPC ditolak dengan pesan error 'Hanya Super Admin' / 'Hanya Admin Sales / Super Admin' (42501).",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-07", "RBAC", "P0", "Operator Manager", "Jakarta",
         "Operator Manager dilarang approve availability host & akses host rates",
         "1. Login sebagai Operator Manager.\n2. Coba panggil action 'review_availability' atau query snapshot rates.",
         "Aksi review ditolak ('Hanya Host Manager') dan array 'rates' selalu kosong ([]).",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-08", "RBAC", "P0", "Host Manager", "Jakarta",
         "Host Manager dilarang create Admin Sales / Super Admin",
         "1. Login sebagai Host Manager.\n2. Buka form Undang Akun di menu Host Cost.",
         "Pilihan role hanya 'Host'. Opsi Host Manager dan Admin Sales tidak tersedia.",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-09", "RBAC", "P0", "Host Manager", "Jakarta",
         "Isolasi lokasi Host Manager (Jakarta vs Bandung)",
         "1. Login sebagai Host Manager Jakarta.\n2. Coba review availability/izin host Bandung via RPC langsung.",
         "Aksi ditolak karena host berada di luar lokasi manager.",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-10", "RBAC", "P0", "Operator Staff", "Jakarta",
         "Operator Staff hanya memiliki akses baca jadwal published & checklist",
         "1. Login sebagai Operator Staff.\n2. Buka Jadwal Livestreaming.\n3. Coba panggil action publish/delete.",
         "Tombol Edit, Delete, Publish All tidak tampil. Sesi berstatus draft tidak terlihat. RPC mutasi ditolak.",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-11", "RBAC", "P0", "Operator Staff / Host", "Jakarta",
         "Staff/Host dilarang membaca data lokasi lain",
         "1. Panggil production_snapshot_v2 dengan p_location='bandung' menggunakan token user Jakarta.",
         "RPC langsung melempar error 'Akses ditolak' (42501).",
         "Not Tested", "", "", "", ""),
        
        ("RBAC-12", "RBAC", "P0", "Super Admin", "All",
         "Super Admin memiliki akses global dan switch lokasi",
         "1. Login sebagai Super Admin.\n2. Beralih antara lokasi Jakarta dan Bandung, serta buka halaman dengan snapshot global.",
         "Super Admin dapat melihat dan mengelola seluruh data quotation, master brand, studio, plotting, dan jadwal di kedua kota.",
         "Not Tested", "", "", "", ""),

        # Domain 3: Master Data (Brand & Studio) (MSTR)
        ("MSTR-01", "Master Data", "P0", "Super Admin", "All",
         "Tambah Master Brand baru dengan Shop ID valid",
         "1. Buka Master Data.\n2. Isi Nama brand: 'Brand Alpha', TikTok: 'TT-ALPHA', Shopee: 'SP-ALPHA', Mirror: 'ST-ALPHA'.\n3. Klik Tambah brand.",
         "Brand berhasil ditambahkan dan langsung muncul pada tabel Master Brand & Shop ID.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-02", "Master Data", "P1", "Super Admin", "All",
         "Validasi prefix format Shop ID Brand",
         "1. Masukkan TikTok dengan prefix 'XX-123' atau Shopee 'SH-123'.\n2. Klik Tambah brand.",
         "Form/server menolak input dengan error pattern validasi (TikTok harus TT%, Shopee SP%, Mirror ST%).",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-03", "Master Data", "P1", "Super Admin", "All",
         "Validasi nama Brand duplikat atau kosong",
         "1. Masukkan nama brand yang sudah ada atau hanya spasi.\n2. Submit form.",
         "Aplikasi menolak pembuatan brand duplikat / kosong dengan pesan validasi.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-04", "Master Data", "P0", "Super Admin", "All",
         "Edit Master Brand yang belum terpakai quotation",
         "1. Klik Edit pada brand yang belum digunakan quotation.\n2. Ubah nama dan Shop ID.\n3. Klik Simpan perubahan brand.",
         "Perubahan tersimpan dengan sukses di database.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-05", "Master Data", "P0", "Super Admin", "All",
         "Edit nama Brand yang aktif di quotation (propagasi nama)",
         "1. Klik Edit pada Brand yang sudah memiliki quotation aktif.\n2. Ubah nama brand (e.g. 'Fonterra' -> 'Fonterra Brands Indonesia').\n3. Simpan.",
         "Nama brand berhasil diubah dan nama pada tabel production_quotations otomatis ikut ter-update.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-06", "Master Data", "P0", "Super Admin", "All",
         "Edit Shop ID pada Brand yang aktif di quotation (dicegah)",
         "1. Klik Edit pada Brand yang sudah memiliki quotation aktif.\n2. Coba ubah Shop ID TikTok/Shopee/Mirror.\n3. Simpan.",
         "Sistem menolak perubahan Shop ID dengan pesan 'Shop ID tidak dapat diubah selama brand masih dipakai quotation'.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-07", "Master Data", "P0", "Super Admin", "All",
         "Hapus Master Brand yang aktif di quotation (dicegah)",
         "1. Klik Delete pada brand yang masih memiliki quotation.\n2. Konfirmasi dialog popup.",
         "Penghapusan ditolak dengan pesan 'Brand masih dipakai quotation. Hapus atau ubah quotation terlebih dahulu.'",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-08", "Master Data", "P2", "Super Admin", "All",
         "Hapus Master Brand yang tidak terpakai",
         "1. Klik Delete pada brand baru yang belum pernah dipakai quotation.\n2. Konfirmasi dialog popup.",
         "Brand berhasil dihapus permanen dari sistem.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-09", "Master Data", "P0", "Super Admin", "Jakarta / Bandung",
         "Tambah Master Studio baru dengan kapasitas valid (1–50)",
         "1. Buka Master Studio & Lokasi.\n2. Pilih Lokasi: 'Jakarta', Nama: 'Studio 1 JKT', Kapasitas: 2.\n3. Klik Tambah studio.",
         "Studio berhasil dibuat dengan 2 slot kapasitas live paralel.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-10", "Master Data", "P1", "Super Admin", "All",
         "Validasi kapasitas studio di luar batas (min 1, max 50)",
         "1. Masukkan kapasitas 0 atau 51.\n2. Submit form.",
         "Input ditolak oleh validasi client dan database constraint.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-11", "Master Data", "P0", "Super Admin", "Jakarta / Bandung",
         "Edit Studio (nama dan ekspansi kapasitas) pada studio berjadwal",
         "1. Klik Edit pada studio yang memiliki jadwal aktif.\n2. Ubah nama dan naikkan kapasitas dari 1 ke 3.\n3. Simpan.",
         "Perubahan nama dan penambahan kapasitas berhasil disimpan.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-12", "Master Data", "P0", "Super Admin", "Jakarta / Bandung",
         "Coba kurangi kapasitas studio di bawah lane terpakai (dicegah)",
         "1. Pada studio yang memiliki jadwal di lane 2, coba edit kapasitas menjadi 1.\n2. Simpan.",
         "Sistem menolak dengan pesan 'Kapasitas baru lebih kecil dari slot jadwal terpakai'.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-13", "Master Data", "P0", "Super Admin", "All",
         "Coba ganti lokasi studio yang memiliki jadwal aktif (dicegah)",
         "1. Coba ganti lokasi studio dari Jakarta ke Bandung pada studio yang memiliki sesi.\n2. Simpan.",
         "Sistem menolak dengan pesan 'Lokasi studio tidak dapat diganti selama masih memiliki jadwal'.",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-14", "Master Data", "P0", "Super Admin", "Jakarta / Bandung",
         "Hapus Studio yang sedang memiliki jadwal aktif (dicegah)",
         "1. Klik Delete pada studio yang memiliki jadwal aktif.\n2. Konfirmasi dialog popup.",
         "Penghapusan ditolak dengan pesan 'Studio masih dipakai jadwal. Pindahkan atau hapus jadwal terlebih dahulu.'",
         "Not Tested", "", "", "", ""),
        
        ("MSTR-15", "Master Data", "P2", "Super Admin", "Jakarta / Bandung",
         "Hapus Studio tanpa jadwal",
         "1. Klik Delete pada studio kosong tanpa jadwal.\n2. Konfirmasi dialog popup.",
         "Studio berhasil dihapus permanen.",
         "Not Tested", "", "", "", ""),

        # Domain 4: Commercial Quotation & Quota Tracker (QUOT)
        ("QUOT-01", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Input Quotation baru dengan parameter valid",
         "1. Buka Quotation & Kuota.\n2. Isi No. Quotation: 'Q-2026-001', Brand: 'Brand Alpha', Platform: 'TikTok', Periode: 2026-10-01 s/d 2026-10-31, Kuota: 50 jam, Rate: 100.000.\n3. Klik Simpan quotation.",
         "Quotation berhasil disimpan. Shop ID otomatis terisi sesuai platform (TT-ALPHA) dan muncul pada Quota Tracker dengan Sisa Kuota 50 jam.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-02", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Resolusi otomatis Shop ID berdasarkan platform Quotation",
         "1. Pilih Platform: 'Shopee' -> Shop ID menjadi 'SP-ALPHA'.\n2. Pilih Platform: 'Mirror' -> Shop ID menjadi 'ST-ALPHA'.",
         "Shop ID otomatis ter-resolve dari Master Brand tanpa perlu input manual.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-03", "Quotation & Kuota", "P1", "Admin Sales", "All",
         "Validasi nomor referensi quotation duplikat",
         "1. Masukkan nomor quotation yang sudah terdaftar sebelumnya.\n2. Submit form.",
         "Sistem menolak dengan error unique constraint violation pada nomor referensi quotation.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-04", "Quotation & Kuota", "P1", "Admin Sales", "All",
         "Validasi tanggal periode quotation (akhir < mulai atau rentang > 366 hari)",
         "1. Masukkan tanggal akhir sebelum tanggal mulai, atau rentang lebih dari 366 hari.\n2. Submit form.",
         "Form dan database constraint menolak periode yang tidak valid.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-05", "Quotation & Kuota", "P1", "Admin Sales", "All",
         "Validasi batas kuota jam quotation (1–10000)",
         "1. Masukkan kuota 0 atau 15000.\n2. Submit form.",
         "Input kuota jam di luar rentang 1–10000 ditolak.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-06", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Perhitungan Quota Tracker otomatis saat jadwal terplot",
         "1. Periksa baris quotation di tabel Quota Tracker.\n2. Buat sesi jadwal 4 jam pada quotation tersebut.\n3. Refresh/buka kembali Quota Tracker.",
         "Terplot bertambah 4 jam, Sisa berkurang 4 jam, Kuota tetap 50 jam.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-07", "Quotation & Kuota", "P2", "Admin Sales", "All",
         "Pembulatan nilai total quotation ke ribuan terdekat",
         "1. Buat quotation dengan rate 83.333 dan kuota 10 jam.\n2. Amati kolom Nilai Quotation pada tabel.",
         "Nilai quotation terhitung dan dibulatkan menjadi Rp 833.000.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-08", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Edit quotation yang belum memiliki jadwal",
         "1. Klik Edit pada quotation tanpa jadwal.\n2. Ubah brand, platform, tanggal periode, kuota, dan rate.\n3. Simpan perubahan.",
         "Semua field quotation berhasil diperbarui.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-09", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Coba ganti Brand/Platform pada quotation yang memiliki jadwal (dicegah)",
         "1. Klik Edit pada quotation yang sudah memiliki jadwal terplot.\n2. Coba ganti Brand atau Platform.\n3. Simpan.",
         "Sistem menolak dengan pesan 'Brand/platform tidak dapat diganti selama quotation masih memiliki jadwal'.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-10", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Coba kurangi kuota jam di bawah total jam terplot (dicegah)",
         "1. Pada quotation dengan 10 jam terplot, coba ubah kuota jam menjadi 5 jam.\n2. Simpan.",
         "Sistem menolak dengan pesan 'Kuota tidak boleh lebih kecil dari jam terplot'.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-11", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Coba persempit tanggal periode hingga mengeluarkan jadwal aktif (dicegah)",
         "1. Quotation memiliki jadwal di tanggal 2026-10-25.\n2. Coba ubah periode akhir menjadi 2026-10-20.\n3. Simpan.",
         "Sistem menolak dengan pesan 'Periode harus mencakup seluruh jadwal aktif'.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-12", "Quotation & Kuota", "P0", "Admin Sales", "All",
         "Hapus quotation dengan cascade delete pada seluruh jadwal & checklist",
         "1. Klik Delete pada quotation.\n2. Baca peringatan konfirmasi penghapusan.\n3. Klik OK.",
         "Quotation, semua sesi jadwal terkait di seluruh tanggal/kota, dan semua checklist terkait terhapus permanen; data logbook ter-update.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-13", "Quotation & Kuota", "P2", "Admin Sales", "All",
         "Batal hapus quotation pada dialog konfirmasi",
         "1. Klik Delete pada quotation.\n2. Klik Cancel pada dialog konfirmasi.",
         "Quotation dan jadwal tetap utuh tanpa perubahan.",
         "Not Tested", "", "", "", ""),
        
        ("QUOT-14", "Quotation & Kuota", "P1", "Operator Manager", "Jakarta",
         "Operator Manager melihat referensi Quota Tracker tanpa aksi edit/delete",
         "1. Login sebagai Operator Manager dan buka Quotation & Kuota.",
         "Tabel Quota Tracker tampil lengkap dengan rate, namun kolom Aksi dan form edit/delete tidak muncul.",
         "Not Tested", "", "", "", ""),

        # Domain 5: Best Hours & Schedule Plotting (PLOT)
        ("PLOT-01", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Atur Best Hours berdurasi pada Quotation",
         "1. Buka menu Plotting Jadwal -> section 1. Atur best hour berdurasi.\n2. Pilih quotation, masukkan rentang: '09-11, 19-22'.\n3. Klik Simpan best hour.",
         "Best hours tersimpan sebagai 2 blok sesi utuh (blok 2 jam dan blok 3 jam) di field best_hour_slots.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-02", "Plotting Jadwal", "P1", "Manager", "Jakarta / Bandung",
         "Validasi format rentang Best Hours",
         "1. Masukkan format salah seperti '11-09', '25-26', atau '09:30-11:00'.\n2. Simpan.",
         "Aplikasi melempar error validasi format best hour yang informatif.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-03", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Plot sesi manual mode 'Tanggal ini saja'",
         "1. Section 2. Plot sesi manual.\n2. Pilih Quotation, Studio, Tanggal, Jam 09:00–11:00, Mode 'Tanggal ini saja'.\n3. Klik Simpan draft.",
         "Satu sesi draft 2 jam berhasil dibuat pada tanggal tersebut dengan slot kapasitas studio otomatis.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-04", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Plot sesi manual mode 'Semua hari di bulan ini'",
         "1. Pilih Quotation dengan kuota cukup.\n2. Pilih Mode 'Semua hari di bulan ini'.\n3. Simpan draft.",
         "Sesi draft otomatis dibuat untuk setiap hari aktif di bulan tersebut yang berada dalam periode quotation.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-05", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Plot sesi manual mode 'Tanggal pilihan' dengan MultiDatePicker",
         "1. Pilih Mode 'Tanggal pilihan'.\n2. Klik beberapa tanggal di kalender (e.g. 5, 8, 12, 19).\n3. Hapus salah satu chip tanggal pilihan.\n4. Klik Simpan draft.",
         "Sesi draft berhasil dibuat hanya pada tanggal-tanggal yang dipilih.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-06", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Deduplikasi tanggal pada manual multi-date plotting",
         "1. Pilih rentang tanggal yang mengandung tanggal duplikat dalam payload.\n2. Simpan draft.",
         "Server mendeduplikasi tanggal sebelum kalkulasi kapasitas; tidak memicu error kapasitas penuh palsu.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-07", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Penentuan otomatis slot lane kapasitas studio",
         "1. Studio memiliki kapasitas 2 live bersamaan.\n2. Buat sesi jam 09–11 (terisi di Lane 1).\n3. Buat sesi kedua jam 09–11 di studio yang sama.",
         "Sesi kedua otomatis dialokasikan ke Lane 2 tanpa bentrok.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-08", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Deteksi bentrok kapasitas studio penuh",
         "1. Studio kapasitas 1 sudah terisi sesi jam 09–11.\n2. Coba buat sesi lain di studio dan jam yang sama.",
         "Sistem menolak dengan pesan error yang menyebutkan nama studio, tanggal, dan jam bentrok.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-09", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Deteksi bentrok akun toko (Shop Account Collision)",
         "1. Akun toko TikTok brand A sudah memiliki sesi live jam 09–11 di Studio 1.\n2. Coba jadwalkan akun toko yang sama pada jam 09–11 di Studio 2.",
         "Sistem menolak dengan pesan error 'Brand/platform sudah memiliki jadwal pada tanggal/jam tersebut'.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-10", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Validasi kuota quotation terlampaui saat manual plotting",
         "1. Sisa kuota quotation tinggal 2 jam.\n2. Coba buat sesi manual 4 jam.",
         "Sistem menolak pembuatan sesi dengan pesan 'Melebihi sisa kuota quotation'.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-11", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Manual plot dengan Host tidak available / izin / bentrok",
         "1. Pilih host yang belum approved availability atau sedang izin pada jam tersebut.\n2. Simpan draft.",
         "Sistem menolak penugasan host dengan pesan 'Host tidak available pada tanggal/jam tersebut'.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-12", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Auto-plot best hours tanpa host (auto_plot_blocks)",
         "1. Section 3. Auto-plot dari best hour tanpa host.\n2. Pilih Quotation dan Studio.\n3. Klik Generate draft.",
         "Sistem otomatis meng-generate sesi-sesi draft sesuai best hours dari tanggal paling awal hingga kuota terpenuhi atau kapasitas habis.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-13", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Plot host otomatis ke jadwal yang sudah ada (auto_assign_hosts)",
         "1. Section 4. Plot host otomatis ke jadwal yang sudah dibuat.\n2. Pilih Host Prioritas 1, 2, 3.\n3. Klik Plot host ke jadwal.",
         "Host yang available otomatis ter-assign ke sesi-sesi unassigned sebagai pending proposal tanpa double-booking.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-14", "Plotting Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Auto-plot best hour + host pilihan sekaligus",
         "1. Section 5. Auto-plot best hour + host pilihan.\n2. Pilih Quotation, Studio, dan Host Prioritas 1–3.\n3. Klik Generate draft.",
         "Sesi draft ter-generate dan langsung di-assign host yang sesuai availability.",
         "Not Tested", "", "", "", ""),
        
        ("PLOT-15", "Plotting Jadwal", "P1", "Manager", "Jakarta / Bandung",
         "Idempotensi auto-plot saat kuota quotation sudah terisi penuh",
         "1. Jalankan auto-plot pada quotation yang kuotanya sudah habis terplot.",
         "Sistem mengembalikan 0 draft dibuat dan 0 sisa tanpa membuat duplikasi.",
         "Not Tested", "", "", "", ""),

        # Domain 6: Studio Timeline & Capacity (TIMEL)
        ("TIMEL-01", "Studio Timeline", "P0", "Manager", "Jakarta / Bandung",
         "Tampilan visual Studio Timeline 24 jam (00:00–23:00)",
         "1. Buka tab Studio Timeline.\n2. Amati baris studio dan kolom jam 00:00 sampai 23:00.",
         "Grid timeline menampilkan seluruh studio dan slot lane kapasitasnya dengan header jam 24 jam lengkap.",
         "Not Tested", "", "", "", ""),
        
        ("TIMEL-02", "Studio Timeline", "P0", "Manager", "Jakarta / Bandung",
         "Indikator status Sesi Draft dan Published pada Timeline",
         "1. Lihat cell studio pada jam yang memiliki sesi draft dan published.",
         "Cell menampilkan nama Brand dengan kelas styling occupied status ('draft' atau 'published').",
         "Not Tested", "", "", "", ""),
        
        ("TIMEL-03", "Studio Timeline", "P2", "Manager", "Jakarta / Bandung",
         "Tooltip detail sesi saat hover cell timeline",
         "1. Hover kursor mouse pada cell yang terisi sesi.",
         "Tooltip menampilkan Live Code, Brand, Platform, dan rentang jam live.",
         "Not Tested", "", "", "", ""),
        
        ("TIMEL-04", "Studio Timeline", "P2", "Manager", "Jakarta / Bandung",
         "Kalkulasi metrik kapasitas studio-jam harian",
         "1. Periksa teks ringkasan di bawah timeline.",
         "Menampilkan total studio-jam terpakai berbanding total kapasitas harian (e.g. 'Terpakai 8 dari 48 studio-jam').",
         "Not Tested", "", "", "", ""),
        
        ("TIMEL-05", "Studio Timeline", "P1", "Manager", "Jakarta / Bandung",
         "Sesi Mirror menggunakan 1 slot kapasitas studio",
         "1. Buat sesi Mirror pada studio tertentu.\n2. Cek timeline.",
         "Sesi Mirror menempati 1 lane slot kapasitas studio.",
         "Not Tested", "", "", "", ""),
        
        ("TIMEL-06", "Studio Timeline", "P1", "Super Admin", "All",
         "Ganti tanggal timeline dan switch lokasi",
         "1. Ubah input tanggal timeline dan dropdown Lokasi Jakarta/Bandung.\n2. Klik Refresh.",
         "Timeline menampilkan studio dan jadwal kota/tanggal yang dipilih.",
         "Not Tested", "", "", "", ""),

        # Domain 7: Schedule Table & Staged Publication (SCHED)
        ("SCHED-01", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Tampilan default Jadwal Livestreaming",
         "1. Buka menu Jadwal Livestreaming.\n2. Amati filter tanggal dan pagination default.",
         "Filter tanggal default Daily (hari ini WIB), baris per halaman 10, dan tombol aksi Refresh, Edit, Publish Host, Publish All, Delete tersusun rapi.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-02", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Filter rentang tanggal (Daily, Weekly, Monthly, Custom Date)",
         "1. Pilih mode 'Monthly' -> memuat sesi dari tanggal 1 sampai akhir bulan (termasuk tanggal lampau).\n2. Pilih mode 'Custom Date' -> masukkan rentang tanggal acak.",
         "Tabel memuat sesi yang berada dalam rentang tanggal yang dipilih.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-03", "Jadwal Livestreaming", "P1", "Semua Role", "Jakarta / Bandung",
         "Filter kombinasi Lokasi, Platform, Brand, dan Host",
         "1. Terapkan filter Platform: 'TikTok', Brand: 'Brand Alpha', Host: 'Nama Host'.",
         "Tabel memfilter sesi yang cocok sebelum pagination diterapkan.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-04", "Jadwal Livestreaming", "P0", "Host", "Jakarta / Bandung",
         "Toggle Host: 'Semua Jadwal' vs 'Jadwal Saya'",
         "1. Login sebagai Host.\n2. Klik 'Semua Jadwal' -> melihat seluruh jadwal published di kotanya.\n3. Klik 'Jadwal Saya' -> hanya melihat sesi terkonfirmasi miliknya.",
         "Toggle berfungsi instan dan menyaring data dengan benar.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-05", "Jadwal Livestreaming", "P0", "Operator Staff", "Jakarta / Bandung",
         "Operator Staff hanya melihat sesi berstatus Published",
         "1. Login sebagai Operator Staff.\n2. Buka Jadwal Livestreaming pada tanggal yang memiliki sesi draft dan published.",
         "Hanya sesi published yang tampil; sesi draft tersembunyi sepenuhnya.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-06", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Publish All tahap awal (Publish Jadwal)",
         "1. Buat beberapa sesi draft tanpa host.\n2. Klik 'Publish All (N)'.\n3. Konfirmasi modal review.",
         "Sesi draft berubah status menjadi 'published' dan visible bagi staff/host. Operator shift & host bersifat opsional pada publish tahap awal.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-07", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Publish Host pada sesi yang sudah published",
         "1. Plot host pada sesi published.\n2. Klik 'Publish Host (N)'.\n3. Konfirmasi.",
         "Penugasan host disahkan (host_published=true), fee host di-snapshot dari host_rates, dan host dapat melihat penugasan di 'Jadwal Saya'.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-08", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Mode Edit inline pada baris tabel jadwal",
         "1. Klik tombol 'Edit'.\n2. Baris jadwal mendatang yang belum check berubah menjadi input editable (Lokasi, Studio, Searchable Host).",
         "Manager dapat mencari host lewat kolom input search dan memilih studio.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-09", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Simpan proposal edit (Save inline) mempertahankan versi published lama",
         "1. Ubah host/studio pada satu baris, klik 'Save'.\n2. Buka akun Host/Staff pada browser lain.",
         "Proposal tersimpan ke live_session_edits (badge 'Perubahan belum dipublish'). Staff/Host tetap melihat versi published lama sebelum dipublish.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-10", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Validasi Dirty State saat edit jadwal",
         "1. Lakukan perubahan pada form baris edit tanpa klik Save.\n2. Coba ganti filter atau pagination.",
         "Sistem mencegah perpindahan dengan pesan error 'Save perubahan baris sebelum mengganti filter'.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-11", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Klik 'Selesai edit' membuka modal PublishReview",
         "1. Setelah menyimpan beberapa proposal baris, klik 'Selesai edit'.",
         "Modal dialog PublishReview muncul menampilkan jumlah jadwal yang akan dipublish, tombol Cancel, dan tombol 'Publish All (N)'.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-12", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Publikasi perubahan atomik & concurrency version check",
         "1. Klik 'Publish All' pada modal PublishReview.\n2. Sistem memvalidasi kapasitas studio, bentrok toko, availability host, dan edit_version.",
         "Perubahan berhasil diterapkan serentak, status has_pending hilang, dan versi baru tampil ke seluruh staff.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-13", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Rollback atomik saat satu sesi dalam batch edit/publish gagal",
         "1. Buat batch publish di mana salah satu host kehilangan availability / rate.\n2. Klik Publish All.",
         "Seluruh batch dibatalkan; tidak ada sesi yang terpublish sebagian; proposal edit tetap tersimpan untuk diperbaiki.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-14", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Mode Delete jadwal (pilih individual / pilih semua hasil filter)",
         "1. Klik tombol 'Delete' -> checkbox muncul.\n2. Klik 'Pilih semua hasil filter' (bisa lintas halaman, max 1000).\n3. Klik 'Hapus terpilih (N)' dan konfirmasi.",
         "Seluruh sesi terpilih dan checklist terkait terhapus permanen; kuota quotation kembali tersedia.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-15", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Pembatalan sesi tunggal (Batal)",
         "1. Klik tombol 'Batal' pada baris sesi draft/published.\n2. Konfirmasi.",
         "Sesi berubah status menjadi 'cancelled', pending edits dihapus, dan kuota quotation dikembalikan.",
         "Not Tested", "", "", "", ""),
        
        ("SCHED-16", "Jadwal Livestreaming", "P0", "Manager", "Jakarta / Bandung",
         "Sesi yang sudah tercatat checklist atau sudah berjalan tidak dapat diedit/dibatalkan",
         "1. Coba klik Batal atau Edit pada sesi lampau yang sudah memiliki live_check.",
         "Sistem menolak dengan pesan 'Sesi sudah berjalan/tercatat; perlu koreksi terpisah'.",
         "Not Tested", "", "", "", ""),

        # Domain 8: Excel Import & Export (EXCEL)
        ("EXCEL-01", "Excel Workflows", "P0", "Manager", "Jakarta / Bandung",
         "Unduh Template Excel Jadwal (livestream-schedule.xlsx)",
         "1. Klik link 'Template Excel' pada toolbar Jadwal Livestreaming.\n2. Buka file hasil unduhan.",
         "File livestream-schedule.xlsx berhasil diunduh, memiliki sheet 'Panduan' dan sheet 'Jadwal' dengan kolom header baku.",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-02", "Excel Workflows", "P0", "Manager", "Jakarta / Bandung",
         "Import file .xlsx valid & pencocokan otomatis quotation",
         "1. Isi sheet Jadwal dengan data valid.\n2. Klik Import Excel dan pilih file.\n3. Amati tabel review import.",
         "File terbaca sukses. Modal review mencocokkan quotation secara otomatis berdasarkan brand, platform, dan rentang tanggal.",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-03", "Excel Workflows", "P0", "Manager", "Jakarta / Bandung",
         "Import dengan multiple kandidat quotation (pemilihan periode)",
         "1. Import jadwal yang brand & platform-nya memiliki lebih dari satu quotation aktif.\n2. Pilih periode yang sesuai pada dropdown per baris di modal review.\n3. Klik Import draft.",
         "Sesi berhasil diimpor sesuai quotation yang dipilih.",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-04", "Excel Workflows", "P1", "Manager", "Jakarta / Bandung",
         "Validasi format jam pada file import (jam bulat wajib)",
         "1. Masukkan jam '9.30' atau '09:30' pada template Excel.\n2. Coba import.",
         "Aplikasi menolak file dengan pesan error spesifik: 'Jam harus bulat, misalnya 9 atau 09:00' tanpa pembulatan sepihak.",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-05", "Excel Workflows", "P1", "Manager", "Jakarta / Bandung",
         "Validasi header file Excel import tidak sesuai",
         "1. Ubah salah satu nama kolom header di sheet Jadwal.\n2. Coba import file.",
         "Import ditolak dengan pesan 'Header tidak sesuai. Unduh Template Excel terbaru, sheet Jadwal, tanpa mengubah urutan kolom.'",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-06", "Excel Workflows", "P1", "Manager", "Jakarta / Bandung",
         "Validasi formula Excel dan error cell pada file import",
         "1. Masukkan formula Excel (e.g. =TODAY()) atau cell error #REF! pada data.\n2. Coba import.",
         "Import ditolak dengan pesan peringatan agar menggunakan nilai biasa, bukan formula.",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-07", "Excel Workflows", "P1", "Manager", "Jakarta / Bandung",
         "Validasi ukuran file > 5 MB atau baris > 500",
         "1. Unggah file .xlsx berukuran > 5 MB atau berisi 501 baris jadwal.",
         "File ditolak dengan pesan error 'Ukuran file maksimal 5 MB' / 'Isi 1–500 baris jadwal per import.'",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-08", "Excel Workflows", "P0", "Manager", "Jakarta / Bandung",
         "Import atomik: rollback seluruh baris saat ada 1 baris invalid",
         "1. Siapkan 10 baris jadwal di mana baris ke-10 kuotanya melebihi sisa quotation.\n2. Lakukan import.",
         "Seluruh import dibatalkan; 9 baris sebelumnya tidak ada yang tersimpan ke database.",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-09", "Excel Workflows", "P0", "Semua Role", "Jakarta / Bandung",
         "Export Jadwal ke .xlsx (mencakup seluruh hasil filter)",
         "1. Terapkan filter tanggal Monthly atau Custom.\n2. Klik 'Export Excel (N)'.",
         "Browser mengunduh file Jadwal_YYYY-MM-DD_YYYY-MM-DD.xlsx berisi seluruh N baris hasil filter di semua halaman pagination.",
         "Not Tested", "", "", "", ""),
        
        ("EXCEL-10", "Excel Workflows", "P0", "Semua Role", "Jakarta / Bandung",
         "Verifikasi struktur dan keamanan data file Export",
         "1. Buka file hasil Export.\n2. Periksa kolom: Tanggal (serial numerik Excel), Brand, Platform, Lokasi, Studio, Jam Mulai, Jam Selesai, Host, Host ID, Status, Session ID, Live ID, Label Jadwal, Legacy Live ID.",
         "Struktur kolom lengkap, cell tanggal numerik valid, nama brand diawali tanda '=' aman dari formula injection, dan kolom rahasia rate/fee TIDAK terekspos.",
         "Not Tested", "", "", "", ""),

        # Domain 9: Schedule Summary / Ringkasan Jadwal (SUMM)
        ("SUMM-01", "Ringkasan Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Tampilan default Ringkasan Jadwal (Monthly)",
         "1. Buka menu Ringkasan Jadwal.\n2. Amati filter tanggal dan header ringkasan.",
         "Default filter mencakup bulan berjalan (1 bulan penuh), menampilkan total jam terjadwal dan total sesi.",
         "Not Tested", "", "", "", ""),
        
        ("SUMM-02", "Ringkasan Jadwal", "P0", "Manager", "Jakarta / Bandung",
         "Pengelompokan total jam berdasarkan Brand, Platform, dan Lokasi",
         "1. Amati tabel Ringkasan Jadwal.",
         "Baris teragregasi berdasarkan kombinasi unik [Brand, Platform, Lokasi] dengan kalkulasi jam terjadwal yang akurat.",
         "Not Tested", "", "", "", ""),
        
        ("SUMM-03", "Ringkasan Jadwal", "P2", "Manager", "Jakarta / Bandung",
         "Sesi berstatus 'cancelled' tidak dihitung dalam ringkasan",
         "1. Batalkan salah satu sesi 2 jam.\n2. Periksa kembali Ringkasan Jadwal.",
         "Total jam terjadwal berkurang 2 jam.",
         "Not Tested", "", "", "", ""),
        
        ("SUMM-04", "Ringkasan Jadwal", "P2", "Manager", "Jakarta / Bandung",
         "Sesi platform Mirror dihitung satu kali sesuai durasi",
         "1. Jadwalkan sesi Mirror 3 jam.\n2. Periksa tabel ringkasan.",
         "Total jam terhitung 3 jam (tidak berlipat ganda).",
         "Not Tested", "", "", "", ""),
        
        ("SUMM-05", "Ringkasan Jadwal", "P2", "Manager", "Jakarta / Bandung",
         "Filter interaktif pada Ringkasan Jadwal",
         "1. Ubah filter Lokasi, Platform, atau Brand pada halaman ringkasan.",
         "Data ringkasan ter-update secara real-time mengikuti filter.",
         "Not Tested", "", "", "", ""),

        # Domain 10: Host Weekly Availability & Review (AVAIL)
        ("AVAIL-01", "Availability Host", "P0", "Host", "Jakarta / Bandung",
         "Tampilan grid availability mingguan Host (7 hari x 24 jam)",
         "1. Login sebagai Host dan buka menu Availability Host.\n2. Amati grid waktu.",
         "Grid menampilkan 7 hari dari tanggal awal filter dengan slot jam 00:00 s/d 23:00.",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-02", "Availability Host", "P0", "Host", "Jakarta / Bandung",
         "Seleksi slot availability via sentuhan / pointer drag",
         "1. Klik dan drag kursor mouse / sentuhan pada rentang jam (e.g. 09:00 s/d 14:00).\n2. Klik ulang untuk membatalkan seleksi.",
         "Slot terpilih berubah warna menjadi 'Dipilih' dan counter tombol update ('Kirim N jam availability').",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-03", "Availability Host", "P0", "Host", "Jakarta / Bandung",
         "Cutoff availability 1 jam sebelum slot waktu",
         "1. Amati slot pada hari ini yang waktunya < 1 jam dari waktu sekarang atau sudah lewat.",
         "Slot berstatus disabled ('—') dan tidak dapat dipilih.",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-04", "Availability Host", "P0", "Host", "Jakarta / Bandung",
         "Submit parsial availability & penguncian status Pending",
         "1. Pilih beberapa slot bebas.\n2. Klik 'Kirim N jam availability'.",
         "Slot berubah warna/status menjadi 'pending' dan terkunci dari perubahan host.",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-05", "Availability Host", "P0", "Host Manager", "Jakarta / Bandung",
         "Host Manager melihat antrean availability pending di kotanya",
         "1. Login sebagai Host Manager dan buka Availability Host.",
         "Daftar per host dengan jumlah jam pending dan rincian tanggal/jam tampil rapi.",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-06", "Availability Host", "P0", "Host Manager", "Jakarta / Bandung",
         "Host Manager menyetujui (Approve) slot availability",
         "1. Klik 'Approve semua dalam rentang' atau tombol Approve pada baris slot.",
         "Status slot berubah menjadi 'approved' dan siap digunakan untuk plotting jadwal.",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-07", "Availability Host", "P0", "Host Manager", "Jakarta / Bandung",
         "Host Manager menolak (Reject) slot availability dengan catatan",
         "1. Klik 'Reject' pada slot pending.\n2. Masukkan alasan penolakan (e.g. 'Revisi jam malam').",
         "Status slot berubah menjadi 'rejected' dan catatan review tersimpan.",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-08", "Availability Host", "P0", "Host", "Jakarta / Bandung",
         "Host mengajukan ulang (resubmit) slot yang berstatus Rejected",
         "1. Login sebagai Host.\n2. Buka grid availability pada tanggal/jam yang sebelumnya di-reject.\n3. Pilih kembali slot tersebut dan klik Kirim.",
         "Slot rejected dapat dipilih dan diajukan kembali menjadi status 'pending'.",
         "Not Tested", "", "", "", ""),
        
        ("AVAIL-09", "Availability Host", "P0", "Operator Manager", "Jakarta / Bandung",
         "Operator Manager dicegah mereview availability Host",
         "1. Login sebagai Operator Manager.\n2. Coba panggil action 'review_availability'.",
         "Aksi ditolak server-side dengan error 'Hanya Host Manager' (42501).",
         "Not Tested", "", "", "", ""),

        # Domain 11: Host Leave / Perizinan (LEAVE)
        ("LEAVE-01", "Perizinan Host", "P0", "Host", "Jakarta / Bandung",
         "Host mengajukan izin pada jadwal published mendatang",
         "1. Login sebagai Host dan buka menu Perizinan Host.\n2. Masukkan tanggal, jam mulai, jam selesai yang overlap dengan jadwal published miliknya, dan alasan izin.\n3. Klik Ajukan izin.",
         "Izin berhasil diajukan dengan status 'pending'.",
         "Not Tested", "", "", "", ""),
        
        ("LEAVE-02", "Perizinan Host", "P1", "Host", "Jakarta / Bandung",
         "Validasi pengajuan izin tanpa jadwal published atau setelah sesi mulai",
         "1. Coba ajukan izin pada tanggal yang tidak memiliki sesi published, atau pada sesi yang sudah mulai.",
         "Sistem menolak dengan pesan error yang sesuai ('Tidak ada jadwal published yang overlap' / 'Izin harus sebelum sesi').",
         "Not Tested", "", "", "", ""),
        
        ("LEAVE-03", "Perizinan Host", "P0", "Host Manager", "Jakarta / Bandung",
         "Host Manager mereview perizinan Host (Approve / Reject)",
         "1. Login sebagai Host Manager dan buka Perizinan Host.\n2. Tinjau permohonan izin pending.",
         "Host Manager dapat melihat alasan izin dan memilih tombol Approve atau Reject dengan catatan.",
         "Not Tested", "", "", "", ""),
        
        ("LEAVE-04", "Perizinan Host", "P0", "Host Manager", "Jakarta / Bandung",
         "Approval izin otomatis melepas Host dari seluruh sesi yang overlap",
         "1. Klik Approve pada permohonan izin.\n2. Konfirmasi dialog 'Approve izin dan lepas host dari sesi yang overlap?'.\n3. Periksa tabel jadwal livestreaming.",
         "Host otomatis terlepas (host_id=null, host_fee=null) dari seluruh sesi yang overlap tanpa perlu membatalkan sesi atau menghapus studio.",
         "Not Tested", "", "", "", ""),
        
        ("LEAVE-05", "Perizinan Host", "P0", "Host Manager", "Jakarta / Bandung",
         "Sesi tetap Published dan Manager dapat mem-plot host pengganti",
         "1. Buka Jadwal Livestreaming pada sesi yang host-nya baru dilepas izin.\n2. Lakukan plot host pengganti yang available.",
         "Sesi tetap berstatus published, reservasi studio utuh, dan host pengganti berhasil di-plot.",
         "Not Tested", "", "", "", ""),
        
        ("LEAVE-06", "Perizinan Host", "P0", "Host Manager", "Jakarta / Bandung",
         "Host yang izinnya disetujui tidak dapat di-plot kembali pada jam izin",
         "1. Coba assign kembali host yang izinnya sudah approved ke sesi pada jam izin tersebut.",
         "Sistem menolak penugasan dengan pesan 'Host tidak available/bentrok/izin'.",
         "Not Tested", "", "", "", ""),

        # Domain 12: Host Cost, Rates & User Invitations (COST)
        ("COST-01", "Host Cost & Rates", "P0", "Host Manager", "Jakarta / Bandung",
         "Tampilan panel accordion Host Cost tertutup secara default",
         "1. Buka menu Host Cost.\n2. Amati section 'Undang akun' dan 'Host & estimasi cost'.",
         "Kedua section tertutup (collapsed) secara default dan dapat dibuka secara independen.",
         "Not Tested", "", "", "", ""),
        
        ("COST-02", "Host Cost & Rates", "P0", "Host Manager", "Jakarta / Bandung",
         "Host Manager mengundang akun Host baru dengan fee awal wajib",
         "1. Buka panel 'Undang akun'.\n2. Masukkan Nama, Email, Role: 'Host', Employment: 'Mitra', Fee awal / jam: 30000, Berlaku mulai: hari ini.\n3. Klik Kirim undangan.",
         "Undangan email terkirim via Netlify Function, profile tercipta di Supabase, dan initial rate tersimpan di host_rates.",
         "Not Tested", "", "", "", ""),
        
        ("COST-03", "Host Cost & Rates", "P0", "Super Admin", "All",
         "Super Admin mengundang Host Manager dan Admin Sales",
         "1. Login sebagai Super Admin.\n2. Buka form Undang Akun di Host Cost.\n3. Pilih Role: 'Host Manager' atau 'Admin Sales'.\n4. Kirim undangan.",
         "Akun berhasil diundang dengan role yang dipilih.",
         "Not Tested", "", "", "", ""),
        
        ("COST-04", "Host Cost & Rates", "P0", "Host Manager", "Jakarta",
         "Host Manager dilarang mengundang akun lokasi lain / eskalasi role",
         "1. Panggil Netlify function accounts dengan payload location='bandung' atau role='super_admin' menggunakan token Host Manager Jakarta.",
         "Request ditolak dengan status HTTP 403 'Akses ditolak' / 'Hanya anggota tim lokasi sendiri'.",
         "Not Tested", "", "", "", ""),
        
        ("COST-05", "Host Cost & Rates", "P0", "Host Manager", "Jakarta / Bandung",
         "Tambah riwayat rate baru untuk Host",
         "1. Buka panel 'Host & estimasi cost' -> 'Tambah rate host'.\n2. Pilih Host, Tanggal berlaku (mendatang), Fee/jam: 35000.\n3. Klik Simpan rate.",
         "Rate baru tersimpan tanpa menimpa rate historis lama.",
         "Not Tested", "", "", "", ""),
        
        ("COST-06", "Host Cost & Rates", "P0", "Host Manager", "Jakarta / Bandung",
         "Validasi rate baru tidak boleh retroaktif sebelum jadwal published",
         "1. Coba masukkan tanggal berlaku rate sebelum jadwal published terakhir dari host tersebut.",
         "Sistem menolak dengan pesan 'Rate harus setelah jadwal published terakhir dan tidak retroaktif'.",
         "Not Tested", "", "", "", ""),
        
        ("COST-07", "Host Cost & Rates", "P0", "Host Manager", "Jakarta / Bandung",
         "Perhitungan estimasi cost terjadwal vs cost aktual di tabel",
         "1. Amati tabel Host & estimasi cost.\n2. Cek kolom: Jam terjadwal, Jam eligible, Cost terjadwal, Cost aktual.",
         "Cost terjadwal = total jam planned * rate snapshot. Cost aktual = total jam actual eligible * rate snapshot.",
         "Not Tested", "", "", "", ""),
        
        ("COST-08", "Host Cost & Rates", "P0", "Host Manager / Host", "Jakarta / Bandung",
         "Export Host Logbook harian ke Excel (.xlsx)",
         "1. Klik tombol 'Logbook harian' pada baris host.\n2. Buka file hasil unduhan.",
         "File Logbook_Host_[Name]_[Start]_[End].xlsx terunduh berisi agregasi durasi aktual harian per tanggal WIB.",
         "Not Tested", "", "", "", ""),

        # Domain 13: Technical & Host Checklists (CHECK)
        ("CHECK-01", "Checklist & Verifikasi", "P0", "Operator Staff", "Jakarta / Bandung",
         "Operator Staff submit Check Operator pada sesi published",
         "1. Login sebagai Operator Staff yang bertugas pada shift sesi live.\n2. Buka Jadwal Livestreaming saat sesi dimulai.\n3. Klik tombol 'Check' -> centang Audio, Kamera, Jaringan -> masukkan catatan opsional -> Kirim check.",
         "Checklist operator tersimpan sukses (status OP di tabel berubah menjadi '✓').",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-02", "Checklist & Verifikasi", "P0", "Operator Staff", "Jakarta / Bandung",
         "Operator yang tidak bertugas pada shift jam sesi ditolak submit check",
         "1. Login sebagai Operator Staff yang tidak memiliki shift assignment pada jam sesi tersebut.\n2. Coba panggil action 'check' kind='operator'.",
         "Ditolak dengan pesan error 'Hanya operator pada shift sesi ini'.",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-03", "Checklist & Verifikasi", "P0", "Operator Staff", "Jakarta / Bandung",
         "Satu check operator meng-cover seluruh sesi live yang overlap dalam shift",
         "1. Ada 2 sesi live paralel di studio berbeda pada jam yang sama.\n2. Satu operator bertugas mengisi check pada salah satu sesi.",
         "Jam shift operator tersebut ter-unlock menjadi eligible di Logbook Production Operator.",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-04", "Checklist & Verifikasi", "P1", "Operator Staff", "Jakarta / Bandung",
         "Pengisian checklist operator duplikat ditolak",
         "1. Coba submit kembali check operator pada sesi yang sudah memiliki check operator.",
         "Sistem menolak dengan error unique constraint (session_id, kind).",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-05", "Checklist & Verifikasi", "P0", "Host", "Jakarta / Bandung",
         "Host yang ditugaskan submit Check Host & durasi aktual",
         "1. Login sebagai Host yang ditugaskan pada sesi published.\n2. Saat sesi dimulai/selesai, klik tombol 'Check'.\n3. Centang Kesiapan, Brief, masukkan jam Mulai aktual & Selesai aktual.\n4. Kirim check.",
         "Check Host tersimpan sukses (status Host di tabel berubah menjadi '✓') dan durasi aktual tercatat.",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-06", "Checklist & Verifikasi", "P0", "Host / Manager", "Jakarta / Bandung",
         "Host lain atau Manager dilarang mengimpersonasi Check Host",
         "1. Coba submit Check Host menggunakan akun host lain atau akun manager.",
         "Sistem menolak dengan pesan 'Hanya Host pada sesi ini'.",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-07", "Checklist & Verifikasi", "P0", "Host", "Jakarta / Bandung",
         "Validasi durasi aktual host (di dalam rentang sesi & tidak di masa depan)",
         "1. Masukkan waktu mulai aktual sebelum sesi dimulai, selesai aktual melewati waktu sekarang, atau selesai <= mulai.\n2. Kirim check.",
         "Sistem menolak dengan pesan 'Durasi aktual harus di dalam sesi dan tidak di masa depan'.",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-08", "Checklist & Verifikasi", "P0", "Host", "Jakarta / Bandung",
         "Penanganan jam selesai aktual pada sesi yang berakhir pukul 24:00",
         "1. Sesi dijadwalkan berakhir pukul 24:00.\n2. Host mengisi jam selesai aktual '00:00'.",
         "Sistem secara cerdas mengonversi '00:00' menjadi waktu tanggal H+1 00:00+07:00 tanpa error.",
         "Not Tested", "", "", "", ""),
        
        ("CHECK-09", "Checklist & Verifikasi", "P1", "Host", "Jakarta / Bandung",
         "Checklist host bersifat immutable setelah terkirim",
         "1. Coba kirim ulang atau edit checklist host yang sudah tersimpan.",
         "Sistem menolak pengiriman ulang (tombol Check hilang / RPC ditolak).",
         "Not Tested", "", "", "", ""),

        # Domain 14: Production Logbooks (LOG)
        ("LOG-01", "Logbook Production", "P0", "Operator Staff", "Jakarta / Bandung",
         "Tampilan Logbook Production Operator (Shift vs Eligible)",
         "1. Login sebagai Operator Staff dan buka menu Logbook Production.\n2. Amati tabel logbook harian.",
         "Tabel menampilkan kolom Tanggal, Jam Shift, Jam Eligible (checklist-gated), dan Pending. Jam shift yang belum terhubung sesi live atau belum dicheck tetap pending.",
         "Not Tested", "", "", "", ""),
        
        ("LOG-02", "Logbook Production", "P0", "Operator Staff", "Jakarta / Bandung",
         "Export Logbook Operator Production ke Excel (.xlsx)",
         "1. Klik tombol 'Export XLSX harian'.\n2. Buka file hasil download.",
         "File Logbook_Production_Operator_[Start]_[End].xlsx terunduh dengan rincian tanggal, nama operator, lokasi, jam eligible, jam shift, dan pending.",
         "Not Tested", "", "", "", ""),
        
        ("LOG-03", "Logbook Production", "P0", "Host", "Jakarta / Bandung",
         "Tampilan Logbook Host summary harian",
         "1. Login sebagai Host dan buka menu Jam & Logbook.\n2. Amati baris ringkasan harian.",
         "Tabel menampilkan Tanggal, Host, Jam Terjadwal, Aktual Eligible, Belum Check, Estimasi Terjadwal, dan Estimasi Aktual.",
         "Not Tested", "", "", "", ""),
        
        ("LOG-04", "Logbook Production", "P0", "Host", "Jakarta / Bandung",
         "Akumulasi durasi aktual harian pada multiple sesi di tanggal yang sama",
         "1. Host memiliki 2 sesi di hari yang sama (Sesi 1: 1.5 jam aktual, Sesi 2: 2.0 jam aktual).\n2. Buka logbook host.",
         "Logbook harian menjumlahkan durasi menjadi 1 baris per tanggal dengan total Aktual Eligible = 3.50 jam.",
         "Not Tested", "", "", "", ""),

        # Domain 15: Live ID & Label System (LIVEID)
        ("LIVEID-01", "Live ID System", "P0", "Semua Role", "All",
         "Generasi otomatis Live Code unik tetap (LS-YY-NNNNNN)",
         "1. Buat sesi livestreaming baru.\n2. Amati kolom Live ID pada tabel jadwal.",
         "Sesi otomatis mendapatkan live_code unik berformat e.g. 'LS-26-000123' yang dihasilkan dari database sequence.",
         "Not Tested", "", "", "", ""),
        
        ("LIVEID-02", "Live ID System", "P0", "Semua Role", "All",
         "Format Label Jadwal dinamis (YYYYMMDD_HHMM-HHMM_SHOPID)",
         "1. Periksa sub-label di bawah Live Code pada tabel jadwal.\n2. Contoh sesi tanggal 2026-09-23 jam 11:00–13:00 dengan Shop ID 'TTFONTERRA'.",
         "Label terformat '20260923_1100-1300_TTFONTERRA' dan otomatis ter-update jika tanggal/jam/akun quotation diubah.",
         "Not Tested", "", "", "", ""),
        
        ("LIVEID-03", "Live ID System", "P2", "Semua Role", "All",
         "Ketersediaan Legacy Live ID (DMYY_HHMM_SHOPID)",
         "1. Hover kursor pada label jadwal atau lakukan Export Excel.",
         "Legacy Live ID alias tersedia e.g. '23926_1113_TTFONTERRA' untuk kebutuhan pelacakan data historis.",
         "Not Tested", "", "", "", ""),
        
        ("LIVEID-04", "Live ID System", "P0", "Semua Role", "All",
         "Stabilitas Live ID saat sesi diedit atau dijadwalkan ulang",
         "1. Ubah jam sesi atau studio sesi melalui inline edit.\n2. Periksa kembali Live ID.",
         "UUID internal dan live_code 'LS-YY-NNNNNN' TETAP SAMA dan tidak berubah.",
         "Not Tested", "", "", "", ""),
        
        ("LIVEID-05", "Live ID System", "P2", "Semua Role", "All",
         "Live ID pada Sesi Mirror",
         "1. Buat sesi live platform Mirror.\n2. Periksa Live ID dan labelnya.",
         "Sesi Mirror memiliki satu UUID dan satu Live Code dengan label mengacu pada Shop ID ST...",
         "Not Tested", "", "", "", ""),

        # Domain 16: Security & Concurrency (SEC)
        ("SEC-01", "Security & Integrity", "P0", "Penetrasi / QA", "All",
         "Row Level Security (RLS) aktif pada semua tabel Production",
         "1. Coba lakukan direct SELECT / INSERT / UPDATE / DELETE via client REST Supabase tanpa RPC pada tabel: production_studios, production_quotations, host_availability, host_leaves, host_rates, live_sessions, live_checks, live_session_edits.",
         "Semua direct query ditolak oleh RLS (permission denied).",
         "Not Tested", "", "", "", ""),
        
        ("SEC-02", "Security & Integrity", "P0", "Penetrasi / QA", "All",
         "Proteksi SQL Injection pada input form text",
         "1. Masukkan payload SQLi (e.g. '; DROP TABLE... --) pada input Nama Brand, Alasan Izin, Catatan Checklist, Studio, dsb.\n2. Submit form.",
         "Input tersimpan aman sebagai string literal murni tanpa dieksekusi oleh query database.",
         "Not Tested", "", "", "", ""),
        
        ("SEC-03", "Security & Integrity", "P0", "Penetrasi / QA", "All",
         "Proteksi XSS (Cross-Site Scripting) pada rendering data",
         "1. Masukkan payload script (e.g. <script>alert(1)</script> atau <img src=x onerror=alert(1)>) pada input catatan atau nama.\n2. Render data di tabel / modal.",
         "String ter-escape dengan benar dan tidak mengeksekusi script apapun di browser.",
         "Not Tested", "", "", "", ""),
        
        ("SEC-04", "Security & Integrity", "P0", "Penetrasi / QA", "All",
         "Proteksi Formula Injection pada Export Excel",
         "1. Buat nama brand yang diawali tanda '=' atau '+' (e.g. '=1+1').\n2. Lakukan Export Excel dan buka di Microsoft Excel.",
         "Cell diexport dengan format text murni / string escaping, tidak dieksekusi sebagai formula di Excel.",
         "Not Tested", "", "", "", ""),
        
        ("SEC-05", "Security & Integrity", "P0", "Netlify / QA", "All",
         "Verifikasi keamanan Netlify Function accounts.ts",
         "1. Coba panggil function tanpa Bearer token, dengan origin liar, atau dengan role yang tidak berwenang.\n2. Uji build Netlify dengan variabel staging.",
         "Function memvalidasi origin, memverifikasi auth token, menolak privilege escalation, dan memblokir lingkungan non-staging.",
         "Not Tested", "", "", "", ""),
        
        ("SEC-06", "Security & Integrity", "P0", "Concurrency / QA", "All",
         "Uji konkurensi: 2 manager mengedit / mempublish jadwal bersamaan",
         "1. Manager A dan Manager B membuka mode edit pada jadwal yang sama.\n2. Manager A menyimpan dan mempublish perubahan.\n3. Manager B mencoba mempublish perubahannya menggunakan versi lama.",
         "Sistem menolak perubahan Manager B dengan pesan 'Perubahan diperbarui manager lain. Refresh dan review ulang.' berkat lock table dan concurrency version check.",
         "Not Tested", "", "", "", ""),

        # Domain 17: Regression Testing (REG)
        ("REG-01", "Regression Testing", "P0", "Super Admin / Manager", "Jakarta / Bandung",
         "Menu Operator Management existing tetap berjalan normal",
         "1. Buka link 'Operator Management ↗' menuju /overview, /availability, /plotting, /cost, /accounts, /logbook.\n2. Cek fungsionalitas shift operator existing.",
         "Seluruh alur existing Operator Management (pengajuan jadwal operator, approval shift, perhitungan cost operator) tetap berfungsi utuh tanpa regresi.",
         "Not Tested", "", "", "", ""),
        
        ("REG-02", "Regression Testing", "P0", "Operator Staff", "Jakarta / Bandung",
         "Beralih antara workspace Operator Management dan Production Ops",
         "1. Login sebagai Operator Staff.\n2. Buka link navigasi antar workspace.",
         "Staff dapat berpindah workspace dengan lancar tanpa kehilangan session.",
         "Not Tested", "", "", "", ""),
        
        ("REG-03", "Regression Testing", "P0", "Operator Manager", "Jakarta / Bandung",
         "Plotting shift operator existing tidak terganggu oleh sesi live",
         "1. Buat dan publish jadwal shift operator mingguan seperti biasa.\n2. Cek apakah shift tersebut otomatis menjadi dasar gating checklist sesi live.",
         "Shift operator berhasil dipublish dan terintegrasi dengan validasi checklist operator live.",
         "Not Tested", "", "", "", ""),
        
        ("REG-04", "Regression Testing", "P0", "Operator Manager", "Jakarta / Bandung",
         "Logbook operator lama tetap independen dan tidak tercampur",
         "1. Lakukan export logbook operator lama di /logbook.\n2. Bandingkan dengan export Logbook Production Operator di /production/logbook.",
         "Kedua export memiliki format terpisah dan data tidak saling menimpa.",
         "Not Tested", "", "", "", "")
    ]

    # -------------------------------------------------------------
    # DATA DEFINITIONS: UAT Mobile Test Cases
    # -------------------------------------------------------------
    mob_cases = [
        ("MOB-01", "Navigasi Mobile", "P0", "Manager & Staff", "Jakarta / Bandung",
         "Drawer Navigasi & Menu Sidebar pada HP",
         "1. Buka aplikasi di browser smartphone.\n2. Buka dan tutup menu drawer.\n3. Beralih antar menu (Ringkasan, Jadwal, Availability, Logbook).",
         "Drawer menu responsif, link menu mudah disentuh, tidak ada tombol yang terpotong atau tumpang tindih.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-02", "Navigasi Mobile", "P0", "Semua Role", "Jakarta / Bandung",
         "Toolbar Filter Periode & Native Date Picker pada HP",
         "1. Buka toolbar filter periode pada HP.\n2. Pilih mode Daily, Weekly, Monthly, Custom Date.\n3. Gunakan native date picker pada smartphone.",
         "Dropdown mode tanggal pas di layar, input native date picker HP terbuka normal, dan data ter-filter sesuai.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-03", "Quotation Mobile", "P0", "Admin Sales", "Jakarta / Bandung",
         "Form Input Quotation & Quota Tracker pada layar sempit",
         "1. Buka Quotation & Kuota pada smartphone portrait.\n2. Isi form input quotation.\n3. Geser horizontal pada tabel Quota Tracker.",
         "Form grid tersusun vertikal yang nyaman diisi, tabel dapat di-scroll horizontal dengan mulus tanpa merusak layout luar.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-04", "Master Data Mobile", "P1", "Super Admin", "Jakarta / Bandung",
         "Master Data Brand & Studio pada tampilan mobile",
         "1. Buka menu Master Data di HP.\n2. Cek form tambah brand/studio dan tabel master.",
         "Form dan tabel master data tertata rapi, tombol aksi Edit/Delete mudah di-tap.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-05", "Plotting Mobile", "P0", "Manager", "Jakarta / Bandung",
         "Form Manual Plotting & MultiDatePicker pada layar sentuh",
         "1. Buka menu Plotting Jadwal di HP.\n2. Pilih Mode 'Tanggal pilihan'.\n3. Tap beberapa tanggal pada komponen MultiDatePicker.\n4. Tap tanda silang (x) pada chip tanggal untuk membatalkan.",
         "Kalender sentuh responsif, chip tanggal terpilih tersusun rapi dan mudah di-tap.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-06", "Studio Timeline Mobile", "P0", "Manager", "Jakarta / Bandung",
         "Studio Timeline 24 jam horizontal scroll & touch tooltip",
         "1. Buka tab Studio Timeline di HP.\n2. Swipe horizontal melintasi kolom jam 00 sampai 23.\n3. Tap pada salah satu cell terisi.",
         "Grid timeline dapat di-swipe mulus, header studio tetap terbaca di sisi kiri, dan detail sesi tampil jelas.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-07", "Jadwal Mobile", "P0", "Manager & Staff", "Jakarta / Bandung",
         "Tabel Jadwal Livestreaming: Sticky Header & Scroll pada HP",
         "1. Buka Jadwal Livestreaming di HP.\n2. Scroll vertikal dan swipe horizontal tabel jadwal.\n3. Tap navigasi pagination (Berikutnya / Sebelumnya).",
         "Header tabel tetap sejajar, teks Live ID dan Brand terbaca jelas, tombol pagination mudah ditekan.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-08", "Jadwal Mobile", "P0", "Manager", "Jakarta / Bandung",
         "Mode Edit Inline Jadwal pada smartphone",
         "1. Tap tombol 'Edit' di HP.\n2. Ketik nama host pada input search host.\n3. Pilih host dari dropdown dan tap 'Save'.",
         "Input pencarian host dan dropdown pas di layar mobile, keyboard virtual tidak menutupi tombol Save.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-09", "Jadwal Mobile", "P0", "Manager", "Jakarta / Bandung",
         "Modal Dialog PublishReview pada layar HP",
         "1. Tap 'Selesai edit' -> modal PublishReview terbuka.\n2. Periksa posisi modal dan tombol aksi.",
         "Modal berada di tengah layar secara responsif, teks jumlah publish terbaca utuh, tombol Cancel dan Publish All mudah ditekan.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-10", "Availability Mobile", "P0", "Host", "Jakarta / Bandung",
         "Grid Availability 7-hari: Tap & Drag sentuhan pada HP",
         "1. Buka Availability Host pada smartphone.\n2. Tap slot jam atau geser jari untuk memilih rentang jam.\n3. Tap tombol 'Kirim N jam availability'.",
         "Seleksi sentuh akurat, warna slot berubah menjadi 'Dipilih', dan pengajuan berhasil dikirim.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-11", "Perizinan Mobile", "P0", "Host", "Jakarta / Bandung",
         "Form Pengajuan Izin Host pada smartphone",
         "1. Buka menu Perizinan Host di HP.\n2. Isi tanggal, jam mulai, jam selesai, dan textarea alasan.",
         "Input tanggal/jam dan textarea alasan responsif saat keyboard virtual aktif.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-12", "Perizinan Mobile", "P0", "Host Manager", "Jakarta / Bandung",
         "Review Permohonan Izin Host pada HP",
         "1. Buka Perizinan Host sebagai Host Manager di HP.\n2. Tap tombol Approve atau Reject.",
         "Daftar izin pending terlihat jelas dan tombol aksi dapat di-tap dengan mudah tanpa salah klik.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-13", "Host Cost Mobile", "P0", "Host Manager", "Jakarta / Bandung",
         "Accordion Host Cost & Form Undang Akun pada HP",
         "1. Buka menu Host Cost di smartphone.\n2. Tap summary 'Undang akun' untuk membuka panel.\n3. Tap summary 'Host & estimasi cost'.",
         "Accordion terbuka/tertutup dengan mulus, form undang akun dan tabel estimasi cost pas dengan lebar layar HP.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-14", "Checklist Mobile", "P0", "Operator Staff", "Jakarta / Bandung",
         "Check Operator Form pada smartphone",
         "1. Tap tombol 'Check' pada sesi published di HP.\n2. Tap checkbox Audio, Kamera, Jaringan.\n3. Masukkan catatan dan tap Kirim check.",
         "Checkbox sentuh memiliki area sentuh (touch target) yang cukup besar dan checklist berhasil dikirim.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-15", "Checklist Mobile", "P0", "Host", "Jakarta / Bandung",
         "Check Host Form & Time Picker pada smartphone",
         "1. Tap tombol 'Check' sebagai Host di HP.\n2. Masukkan jam mulai & selesai aktual via native time picker HP.\n3. Tap Kirim check.",
         "Native time picker smartphone terbuka dengan lancar dan data checklist host tersimpan valid.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-16", "Logbook Mobile", "P0", "Staff & Host", "Jakarta / Bandung",
         "Tabel Logbook & Tombol Export pada HP",
         "1. Buka Logbook Production Operator atau Jam & Logbook Host di smartphone.\n2. Tap tombol 'Export XLSX'.",
         "Tabel logbook rapi dan file Excel berhasil diunduh ke penyimpanan smartphone.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-17", "Tampilan Mobile", "P2", "Semua Role", "Jakarta / Bandung",
         "Rotasi Layar HP (Portrait <-> Landscape)",
         "1. Putar smartphone dari posisi Portrait ke Landscape saat membuka tabel jadwal atau form plotting.\n2. Putar kembali ke Portrait.",
         "Layout menyesuaikan dengan lancar, teks tetap terbaca, dan state form tidak ter-reset.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", ""),
        
        ("MOB-18", "Tampilan Mobile", "P2", "Semua Role", "Jakarta / Bandung",
         "Aktivitas Virtual Keyboard pada Form Input",
         "1. Tap pada berbagai field input (search host, catatan, alasan, dsb).\n2. Amati pergerakan layout saat virtual soft keyboard muncul.",
         "Form tidak tertutup oleh keyboard virtual dan viewport kembali normal saat keyboard ditutup.",
         "Android Chrome / iPhone Safari", "Not Tested", "", "", "", "")
    ]

    # Populate Sheet 2: UAT Testing
    for r_idx, row_data in enumerate(uat_cases, 8):
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(row_data, 1):
            cell = ws_uat.cell(r_idx, c_idx, val)
            cell.font = f_cell_bold if c_idx in [1, 3] else f_cell
            cell.fill = fill_curr
            cell.border = border_all
            
            if c_idx in [1, 3, 5, 9, 12]:
                cell.alignment = al_center_top
            elif c_idx == 4:
                cell.alignment = al_left_top
            else:
                cell.alignment = al_left_top
        ws_uat.row_dimensions[r_idx].height = 48
    
    max_uat_row = 7 + len(uat_cases)

    # Populate Sheet 3: UAT Mobile
    for r_idx, row_data in enumerate(mob_cases, 8):
        fill_curr = fill_zebra if r_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(row_data, 1):
            cell = ws_mobile.cell(r_idx, c_idx, val)
            cell.font = f_cell_bold if c_idx in [1, 3] else f_cell
            cell.fill = fill_curr
            cell.border = border_all
            
            if c_idx in [1, 3, 5, 9, 10, 13]:
                cell.alignment = al_center_top
            elif c_idx == 4:
                cell.alignment = al_left_top
            else:
                cell.alignment = al_left_top
        ws_mobile.row_dimensions[r_idx].height = 48
        
    max_mob_row = 7 + len(mob_cases)

    # -------------------------------------------------------------
    # POPULATE SHEET 1: QA Summary
    # -------------------------------------------------------------
    ws_summary.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws_summary['A1'] = "Orbiz Operator Management — QA Testing Summary Report"
    ws_summary['A1'].font = Font(name=FONT_FAMILY, size=16, bold=True, color="1E293B")
    ws_summary['A2'] = "Laporan Eksekutif & Summary Hasil Pengujian Fitur Livestreaming Production V1"
    ws_summary['A2'].font = f_subtitle
    
    # Metadata Card (Rows 4-7)
    meta_info = [
        ("Project", "Orbiz Operator Management System", "QA Lead / Analyst", "Senior QA Engineer & Automation Analyst"),
        ("Target Branch", "feature/livestreaming-production-v1", "Environment", "Netlify Deploy Preview + Staging Supabase"),
        ("Feature Scope", "Livestreaming Production Ops V1 (Quotations, Master Data, Plotting, Staged Publication, Host Cost, Checklists, Exports)", "Timezone / Reference", "Asia/Jakarta (WIB) · Whole-Hour Planning"),
        ("Document Status", "Ready for Manual UAT / Regression Testing", "Live ID Convention", "LS-YY-NNNNNN (Immutable Sequence)")
    ]
    
    for r_offset, (k1, v1, k2, v2) in enumerate(meta_info, 4):
        ws_summary.cell(r_offset, 1, k1).font = f_kpi_label
        ws_summary.cell(r_offset, 1).fill = fill_kpi
        ws_summary.cell(r_offset, 1).border = border_all
        ws_summary.cell(r_offset, 1).alignment = al_left_center
        
        ws_summary.cell(r_offset, 2, v1).font = f_cell_bold
        ws_summary.cell(r_offset, 2).fill = fill_kpi
        ws_summary.cell(r_offset, 2).border = border_all
        ws_summary.cell(r_offset, 2).alignment = al_left_center
        
        ws_summary.cell(r_offset, 4, k2).font = f_kpi_label
        ws_summary.cell(r_offset, 4).fill = fill_kpi
        ws_summary.cell(r_offset, 4).border = border_all
        ws_summary.cell(r_offset, 4).alignment = al_left_center
        
        ws_summary.cell(r_offset, 5, v2).font = f_cell_bold
        ws_summary.cell(r_offset, 5).fill = fill_kpi
        ws_summary.cell(r_offset, 5).border = border_all
        ws_summary.cell(r_offset, 5).alignment = al_left_center
        
        ws_summary.row_dimensions[r_offset].height = 20

    # Section 1: Executive KPI Metrics Cards (Row 9-11)
    ws_summary['A9'] = "1. EXECUTIVE TEST METRICS"
    ws_summary['A9'].font = f_sec_header
    
    kpi_cards = [
        ("A10", "A11", "TOTAL TEST CASES", f"=COUNTA('UAT Testing'!A8:A{max_uat_row})+COUNTA('UAT Mobile'!A8:A{max_mob_row})"),
        ("B10", "B11", "CRITICAL (P0)", f"=COUNTIF('UAT Testing'!C8:C{max_uat_row},\"P0\")+COUNTIF('UAT Mobile'!C8:C{max_mob_row},\"P0\")"),
        ("C10", "C11", "HIGH (P1)", f"=COUNTIF('UAT Testing'!C8:C{max_uat_row},\"P1\")+COUNTIF('UAT Mobile'!C8:C{max_mob_row},\"P1\")"),
        ("D10", "D11", "MEDIUM (P2)", f"=COUNTIF('UAT Testing'!C8:C{max_uat_row},\"P2\")+COUNTIF('UAT Mobile'!C8:C{max_mob_row},\"P2\")"),
        ("E10", "E11", "PASSED", f"=COUNTIF('UAT Testing'!I8:I{max_uat_row},\"Pass\")+COUNTIF('UAT Mobile'!J8:J{max_mob_row},\"Pass\")"),
        ("F10", "F11", "FAILED", f"=COUNTIF('UAT Testing'!I8:I{max_uat_row},\"Fail\")+COUNTIF('UAT Mobile'!J8:J{max_mob_row},\"Fail\")"),
        ("G10", "G11", "BLOCKED", f"=COUNTIF('UAT Testing'!I8:I{max_uat_row},\"Blocked\")+COUNTIF('UAT Mobile'!J8:J{max_mob_row},\"Blocked\")"),
        ("H10", "H11", "RETEST", f"=COUNTIF('UAT Testing'!I8:I{max_uat_row},\"Retest\")+COUNTIF('UAT Mobile'!J8:J{max_mob_row},\"Retest\")"),
        ("I10", "I11", "NOT TESTED", f"=COUNTIF('UAT Testing'!I8:I{max_uat_row},\"Not Tested\")+COUNTIF('UAT Mobile'!J8:J{max_mob_row},\"Not Tested\")"),
        ("J10", "J11", "PROGRESS", f"=IFERROR((E11+F11+G11+H11)/A11, 0)"),
    ]
    
    for lbl_c, val_c, title_text, formula_text in kpi_cards:
        c_lbl = ws_summary[lbl_c]
        c_lbl.value = title_text
        c_lbl.font = f_kpi_label
        c_lbl.fill = fill_header
        c_lbl.alignment = al_center
        c_lbl.border = border_all
        
        c_val = ws_summary[val_c]
        c_val.value = formula_text
        c_val.font = f_kpi_val
        c_val.fill = fill_kpi
        c_val.alignment = al_center
        c_val.border = border_all
        if lbl_c == "J10":
            c_val.number_format = '0.0%'
            
    ws_summary.row_dimensions[10].height = 22
    ws_summary.row_dimensions[11].height = 28

    # Section 2: Coverage Breakdown by Domain (Rows 13-33)
    ws_summary['A13'] = "2. DOMAIN TEST COVERAGE BREAKDOWN"
    ws_summary['A13'].font = f_sec_header
    
    domain_headers = [
        "Domain Name", "Domain Scope / Feature Area", 
        "P0", "P1", "P2", "Total", "Pass", "Fail", "Blocked", "Not Tested", "Progress %"
    ]
    for c_idx, h in enumerate(domain_headers, 1):
        cell = ws_summary.cell(14, c_idx, h)
        cell.font = f_tbl_header
        cell.fill = fill_header
        cell.alignment = al_center
        cell.border = border_tbl_header
    ws_summary.row_dimensions[14].height = 26
    
    domains_list = [
        ("Authentication", "Login, session persistence, role landing routes, inactive & profile guard"),
        ("RBAC", "Role-based access control, location data isolation (Jakarta/Bandung), RPC security"),
        ("Master Data", "Master Brand & Shop IDs (TT/SP/ST), Master Studio & kapasitas live paralel"),
        ("Quotation & Kuota", "Input quotation, Quota Tracker, Shop ID mapping, cascade delete, rounding"),
        ("Plotting Jadwal", "Best hours berdurasi, manual single/month/selected plotting, auto-plot blocks"),
        ("Studio Timeline", "Visual timeline 24 jam, kapasitas occupied, tooltip sesi, mirror duration"),
        ("Jadwal Livestreaming", "Tabel jadwal, staged publication, inline edit host/studio, filter & pagination"),
        ("Excel Workflows", "Import Excel template (.xlsx), review quotation matching, export schedule"),
        ("Ringkasan Jadwal", "Monthly summary grouping by Brand/Platform/Location, mirror single count"),
        ("Availability Host", "Grid availability mingguan 7-hari x 24 jam, drag selection, 1-hour cutoff, approval"),
        ("Perizinan Host", "Pengajuan izin host sebelum live, auto unassign host dari overlapping sessions"),
        ("Host Cost & Rates", "Accordion layout, undang akun host via Netlify, rate history, estimasi cost"),
        ("Checklist & Verifikasi", "Check Operator shift-gated, Check Host durasi aktual & midnight 24:00 conversion"),
        ("Logbook Production", "Logbook Operator checklist-gated, Logbook Host durasi harian, export XLSX"),
        ("Live ID System", "Sequential Live Code (LS-YY-NNNNNN), Dynamic Live Label, Legacy DMYY alias"),
        ("Security & Integrity", "RLS protection, SQLi & XSS protection, Excel formula escaping, concurrency lock"),
        ("Regression Testing", "Operator Management existing routes, plotting operator, legacy logbook")
    ]
    
    for idx, (dom_name, dom_desc) in enumerate(domains_list, 15):
        fill_row = fill_zebra if idx % 2 == 0 else fill_white
        ws_summary.cell(idx, 1, dom_name).font = f_cell_bold
        ws_summary.cell(idx, 1).fill = fill_row
        ws_summary.cell(idx, 1).border = border_all
        ws_summary.cell(idx, 1).alignment = al_left_center
        
        ws_summary.cell(idx, 2, dom_desc).font = f_cell
        ws_summary.cell(idx, 2).fill = fill_row
        ws_summary.cell(idx, 2).border = border_all
        ws_summary.cell(idx, 2).alignment = al_left_center
        
        # P0
        ws_summary.cell(idx, 3, f'=COUNTIFS(\'UAT Testing\'!$B$8:$B${max_uat_row}, $A{idx}, \'UAT Testing\'!$C$8:$C${max_uat_row}, "P0")').font = f_cell
        ws_summary.cell(idx, 3).fill = fill_row
        ws_summary.cell(idx, 3).border = border_all
        ws_summary.cell(idx, 3).alignment = al_center
        
        # P1
        ws_summary.cell(idx, 4, f'=COUNTIFS(\'UAT Testing\'!$B$8:$B${max_uat_row}, $A{idx}, \'UAT Testing\'!$C$8:$C${max_uat_row}, "P1")').font = f_cell
        ws_summary.cell(idx, 4).fill = fill_row
        ws_summary.cell(idx, 4).border = border_all
        ws_summary.cell(idx, 4).alignment = al_center
        
        # P2
        ws_summary.cell(idx, 5, f'=COUNTIFS(\'UAT Testing\'!$B$8:$B${max_uat_row}, $A{idx}, \'UAT Testing\'!$C$8:$C${max_uat_row}, "P2")').font = f_cell
        ws_summary.cell(idx, 5).fill = fill_row
        ws_summary.cell(idx, 5).border = border_all
        ws_summary.cell(idx, 5).alignment = al_center
        
        # Total
        ws_summary.cell(idx, 6, f'=SUM(C{idx}:E{idx})').font = f_cell_bold
        ws_summary.cell(idx, 6).fill = fill_row
        ws_summary.cell(idx, 6).border = border_all
        ws_summary.cell(idx, 6).alignment = al_center
        
        # Pass
        ws_summary.cell(idx, 7, f'=COUNTIFS(\'UAT Testing\'!$B$8:$B${max_uat_row}, $A{idx}, \'UAT Testing\'!$I$8:$I${max_uat_row}, "Pass")').font = f_cell
        ws_summary.cell(idx, 7).fill = fill_row
        ws_summary.cell(idx, 7).border = border_all
        ws_summary.cell(idx, 7).alignment = al_center
        
        # Fail
        ws_summary.cell(idx, 8, f'=COUNTIFS(\'UAT Testing\'!$B$8:$B${max_uat_row}, $A{idx}, \'UAT Testing\'!$I$8:$I${max_uat_row}, "Fail")').font = f_cell
        ws_summary.cell(idx, 8).fill = fill_row
        ws_summary.cell(idx, 8).border = border_all
        ws_summary.cell(idx, 8).alignment = al_center
        
        # Blocked
        ws_summary.cell(idx, 9, f'=COUNTIFS(\'UAT Testing\'!$B$8:$B${max_uat_row}, $A{idx}, \'UAT Testing\'!$I$8:$I${max_uat_row}, "Blocked")').font = f_cell
        ws_summary.cell(idx, 9).fill = fill_row
        ws_summary.cell(idx, 9).border = border_all
        ws_summary.cell(idx, 9).alignment = al_center
        
        # Not Tested
        ws_summary.cell(idx, 10, f'=COUNTIFS(\'UAT Testing\'!$B$8:$B${max_uat_row}, $A{idx}, \'UAT Testing\'!$I$8:$I${max_uat_row}, "Not Tested")').font = f_cell
        ws_summary.cell(idx, 10).fill = fill_row
        ws_summary.cell(idx, 10).border = border_all
        ws_summary.cell(idx, 10).alignment = al_center
        
        # Progress %
        ws_summary.cell(idx, 11, f'=IFERROR((G{idx}+H{idx}+I{idx})/F{idx}, 0)').font = f_cell_bold
        ws_summary.cell(idx, 11).fill = fill_row
        ws_summary.cell(idx, 11).border = border_all
        ws_summary.cell(idx, 11).alignment = al_center
        ws_summary.cell(idx, 11).number_format = '0.0%'
        
        ws_summary.row_dimensions[idx].height = 20

    # Total Row for Domain Table
    tot_row = 15 + len(domains_list)
    ws_summary.cell(tot_row, 1, "TOTAL DESKTOP / WEB UAT").font = f_cell_bold
    ws_summary.cell(tot_row, 1).fill = fill_header_sub
    ws_summary.cell(tot_row, 1).border = border_total
    ws_summary.cell(tot_row, 1).alignment = al_left_center
    
    ws_summary.cell(tot_row, 2, "Semua Domain & Modul Livestreaming Production V1").font = f_subtitle
    ws_summary.cell(tot_row, 2).fill = fill_header_sub
    ws_summary.cell(tot_row, 2).border = border_total
    
    for c_i, col_letter in enumerate(["C", "D", "E", "F", "G", "H", "I", "J"], 3):
        cell = ws_summary.cell(tot_row, c_i, f'=SUM({col_letter}15:{col_letter}{tot_row-1})')
        cell.font = f_kpi_val if c_i == 6 else f_cell_bold
        cell.fill = fill_header_sub
        cell.border = border_total
        cell.alignment = al_center
        
    cell_tot_prog = ws_summary.cell(tot_row, 11, f'=IFERROR((G{tot_row}+H{tot_row}+I{tot_row})/F{tot_row}, 0)')
    cell_tot_prog.font = f_kpi_val
    cell_tot_prog.fill = fill_header_sub
    cell_tot_prog.border = border_total
    cell_tot_prog.alignment = al_center
    cell_tot_prog.number_format = '0.0%'
    ws_summary.row_dimensions[tot_row].height = 24

    # Section 3: Mobile UAT Summary Row
    mob_summary_row = tot_row + 2
    ws_summary.cell(mob_summary_row, 1, "3. MOBILE & RESPONSIVE UAT").font = f_sec_header
    
    mob_tbl_header_row = mob_summary_row + 1
    for c_idx, h in enumerate(["Mobile Suite", "Scope / Device Target", "P0", "P1", "P2", "Total", "Pass", "Fail", "Blocked", "Not Tested", "Progress %"], 1):
        c = ws_summary.cell(mob_tbl_header_row, c_idx, h)
        c.font = f_tbl_header
        c.fill = fill_header
        c.alignment = al_center
        c.border = border_tbl_header
        
    mob_data_row = mob_tbl_header_row + 1
    ws_summary.cell(mob_data_row, 1, "UAT Mobile / Responsive").font = f_cell_bold
    ws_summary.cell(mob_data_row, 1).fill = fill_white
    ws_summary.cell(mob_data_row, 1).border = border_total
    
    ws_summary.cell(mob_data_row, 2, "Smartphone Touch UX (Android Chrome & iPhone Safari)").font = f_cell
    ws_summary.cell(mob_data_row, 2).fill = fill_white
    ws_summary.cell(mob_data_row, 2).border = border_total
    
    ws_summary.cell(mob_data_row, 3, f'=COUNTIF(\'UAT Mobile\'!$C$8:$C${max_mob_row}, "P0")').font = f_cell
    ws_summary.cell(mob_data_row, 3).fill = fill_white
    ws_summary.cell(mob_data_row, 3).border = border_total
    ws_summary.cell(mob_data_row, 3).alignment = al_center
    
    ws_summary.cell(mob_data_row, 4, f'=COUNTIF(\'UAT Mobile\'!$C$8:$C${max_mob_row}, "P1")').font = f_cell
    ws_summary.cell(mob_data_row, 4).fill = fill_white
    ws_summary.cell(mob_data_row, 4).border = border_total
    ws_summary.cell(mob_data_row, 4).alignment = al_center
    
    ws_summary.cell(mob_data_row, 5, f'=COUNTIF(\'UAT Mobile\'!$C$8:$C${max_mob_row}, "P2")').font = f_cell
    ws_summary.cell(mob_data_row, 5).fill = fill_white
    ws_summary.cell(mob_data_row, 5).border = border_total
    ws_summary.cell(mob_data_row, 5).alignment = al_center
    
    ws_summary.cell(mob_data_row, 6, f'=SUM(C{mob_data_row}:E{mob_data_row})').font = f_cell_bold
    ws_summary.cell(mob_data_row, 6).fill = fill_white
    ws_summary.cell(mob_data_row, 6).border = border_total
    ws_summary.cell(mob_data_row, 6).alignment = al_center
    
    ws_summary.cell(mob_data_row, 7, f'=COUNTIF(\'UAT Mobile\'!$J$8:$J${max_mob_row}, "Pass")').font = f_cell
    ws_summary.cell(mob_data_row, 7).fill = fill_white
    ws_summary.cell(mob_data_row, 7).border = border_total
    ws_summary.cell(mob_data_row, 7).alignment = al_center
    
    ws_summary.cell(mob_data_row, 8, f'=COUNTIF(\'UAT Mobile\'!$J$8:$J${max_mob_row}, "Fail")').font = f_cell
    ws_summary.cell(mob_data_row, 8).fill = fill_white
    ws_summary.cell(mob_data_row, 8).border = border_total
    ws_summary.cell(mob_data_row, 8).alignment = al_center
    
    ws_summary.cell(mob_data_row, 9, f'=COUNTIF(\'UAT Mobile\'!$J$8:$J${max_mob_row}, "Blocked")').font = f_cell
    ws_summary.cell(mob_data_row, 9).fill = fill_white
    ws_summary.cell(mob_data_row, 9).border = border_total
    ws_summary.cell(mob_data_row, 9).alignment = al_center
    
    ws_summary.cell(mob_data_row, 10, f'=COUNTIF(\'UAT Mobile\'!$J$8:$J${max_mob_row}, "Not Tested")').font = f_cell
    ws_summary.cell(mob_data_row, 10).fill = fill_white
    ws_summary.cell(mob_data_row, 10).border = border_total
    ws_summary.cell(mob_data_row, 10).alignment = al_center
    
    cell_mob_prog = ws_summary.cell(mob_data_row, 11, f'=IFERROR((G{mob_data_row}+H{mob_data_row}+I{mob_data_row})/F{mob_data_row}, 0)')
    cell_mob_prog.font = f_kpi_val
    cell_mob_prog.fill = fill_white
    cell_mob_prog.border = border_total
    cell_mob_prog.alignment = al_center
    cell_mob_prog.number_format = '0.0%'
    
    ws_summary.row_dimensions[mob_tbl_header_row].height = 24
    ws_summary.row_dimensions[mob_data_row].height = 24

    # Section 4: Testing Guidelines & Defect Standards (Row ~38+)
    guide_start = mob_data_row + 2
    ws_summary.cell(guide_start, 1, "4. PANDUAN EKSEKUSI & STANDAR DEFECT TRACEABILITY").font = f_sec_header
    
    guidelines = [
        ("Status Testing", "Arti / Definisi", "Priority", "Kriteria Penentuan Severity"),
        ("Not Tested", "Test case belum dieksekusi oleh tester (default seluruh test case baru).", "P0 (Critical)", "Fungsi utama gagal total, data integrity rusak, celah keamanan, atau blocker flow."),
        ("Pass", "Hasil aktual sesuai dengan expected result dan didukung evidence screenshot/log valid.", "P1 (Important)", "Fungsi utama dapat berjalan namun terdapat issue fungsional penting atau validasi bocor."),
        ("Fail", "Hasil aktual tidak sesuai expected result atau terjadi crash/error tak terduga.", "P2 (Minor)", "Issue kosmetik minor, alignment teks, atau edge-case sekunder yang tidak mengganggu alur."),
        ("Blocked", "Pengujian terhambat oleh dependency modul lain, bug P0 sebelumnya, atau environment.", "Format Notes Bug", "Gunakan 'Bug: [deskripsi]' jika terkonfirmasi, atau 'Potential Bug: [deskripsi]'."),
        ("Retest", "Test case yang sebelumnya Fail/Blocked telah diperbaiki dan siap diuji ulang.", "Requirement Clarification", "Gunakan jika spesifikasi bisnis membutuhkan konfirmasi lebih lanjut dari Product Owner.")
    ]
    
    for r_i, (c1, c2, c3, c4) in enumerate(guidelines, guide_start + 1):
        f_style = f_tbl_header if r_i == guide_start + 1 else f_cell
        fill_style = fill_header if r_i == guide_start + 1 else fill_white
        
        ws_summary.cell(r_i, 1, c1).font = f_style
        ws_summary.cell(r_i, 1).fill = fill_style
        ws_summary.cell(r_i, 1).border = border_all
        
        ws_summary.cell(r_i, 2, c2).font = f_style
        ws_summary.cell(r_i, 2).fill = fill_style
        ws_summary.cell(r_i, 2).border = border_all
        
        ws_summary.cell(r_i, 4, c3).font = f_style
        ws_summary.cell(r_i, 4).fill = fill_style
        ws_summary.cell(r_i, 4).border = border_all
        
        ws_summary.cell(r_i, 5, c4).font = f_style
        ws_summary.cell(r_i, 5).fill = fill_style
        ws_summary.cell(r_i, 5).border = border_all
        
        ws_summary.row_dimensions[r_i].height = 20

    # -------------------------------------------------------------
    # DATA VALIDATION & CONDITIONAL FORMATTING
    # -------------------------------------------------------------
    # Status Validation: Not Tested,Pass,Fail,Blocked,Retest
    dv_status = DataValidation(type="list", formula1='"Not Tested,Pass,Fail,Blocked,Retest"', allow_blank=True)
    ws_uat.add_data_validation(dv_status)
    dv_status.add(f"I8:I{max_uat_row+50}")
    
    dv_status_mob = DataValidation(type="list", formula1='"Not Tested,Pass,Fail,Blocked,Retest"', allow_blank=True)
    ws_mobile.add_data_validation(dv_status_mob)
    dv_status_mob.add(f"J8:J{max_mob_row+50}")
    
    # Priority Validation: P0,P1,P2
    dv_priority = DataValidation(type="list", formula1='"P0,P1,P2"', allow_blank=True)
    ws_uat.add_data_validation(dv_priority)
    dv_priority.add(f"C8:C{max_uat_row+50}")
    
    dv_priority_mob = DataValidation(type="list", formula1='"P0,P1,P2"', allow_blank=True)
    ws_mobile.add_data_validation(dv_priority_mob)
    dv_priority_mob.add(f"C8:C{max_mob_row+50}")

    # Conditional Formatting Rules for UAT Testing
    rule_pass = CellIsRule(operator='equal', formula=['"Pass"'], fill=fill_pass, font=font_pass)
    rule_fail = CellIsRule(operator='equal', formula=['"Fail"'], fill=fill_fail, font=font_fail)
    rule_blocked = CellIsRule(operator='equal', formula=['"Blocked"'], fill=fill_blocked, font=font_blocked)
    rule_retest = CellIsRule(operator='equal', formula=['"Retest"'], fill=fill_retest, font=font_retest)
    rule_not_tested = CellIsRule(operator='equal', formula=['"Not Tested"'], fill=fill_not_tested, font=font_not_tested)
    
    rule_p0 = CellIsRule(operator='equal', formula=['"P0"'], fill=fill_fail, font=font_fail)
    rule_p1 = CellIsRule(operator='equal', formula=['"P1"'], fill=fill_blocked, font=font_blocked)
    rule_p2 = CellIsRule(operator='equal', formula=['"P2"'], fill=PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"), font=Font(name=FONT_FAMILY, size=9.5, bold=True, color="475569"))

    ws_uat.conditional_formatting.add(f"I8:I{max_uat_row}", rule_pass)
    ws_uat.conditional_formatting.add(f"I8:I{max_uat_row}", rule_fail)
    ws_uat.conditional_formatting.add(f"I8:I{max_uat_row}", rule_blocked)
    ws_uat.conditional_formatting.add(f"I8:I{max_uat_row}", rule_retest)
    ws_uat.conditional_formatting.add(f"I8:I{max_uat_row}", rule_not_tested)
    
    ws_uat.conditional_formatting.add(f"C8:C{max_uat_row}", rule_p0)
    ws_uat.conditional_formatting.add(f"C8:C{max_uat_row}", rule_p1)
    ws_uat.conditional_formatting.add(f"C8:C{max_uat_row}", rule_p2)

    ws_mobile.conditional_formatting.add(f"J8:J{max_mob_row}", rule_pass)
    ws_mobile.conditional_formatting.add(f"J8:J{max_mob_row}", rule_fail)
    ws_mobile.conditional_formatting.add(f"J8:J{max_mob_row}", rule_blocked)
    ws_mobile.conditional_formatting.add(f"J8:J{max_mob_row}", rule_retest)
    ws_mobile.conditional_formatting.add(f"J8:J{max_mob_row}", rule_not_tested)
    
    ws_mobile.conditional_formatting.add(f"C8:C{max_mob_row}", rule_p0)
    ws_mobile.conditional_formatting.add(f"C8:C{max_mob_row}", rule_p1)
    ws_mobile.conditional_formatting.add(f"C8:C{max_mob_row}", rule_p2)

    # -------------------------------------------------------------
    # FREEZE PANES & AUTO FILTERS
    # -------------------------------------------------------------
    ws_uat.freeze_panes = 'A8'
    ws_uat.auto_filter.ref = f"A7:M{max_uat_row}"
    
    ws_mobile.freeze_panes = 'A8'
    ws_mobile.auto_filter.ref = f"A7:N{max_mob_row}"

    # -------------------------------------------------------------
    # COLUMN WIDTHS
    # -------------------------------------------------------------
    summary_col_widths = {
        'A': 24, 'B': 52, 'C': 10, 'D': 10, 'E': 10, 'F': 12, 'G': 10, 'H': 10, 'I': 10, 'J': 14, 'K': 14
    }
    for col_l, width in summary_col_widths.items():
        ws_summary.column_dimensions[col_l].width = width
        
    uat_col_widths = {
        'A': 12,  # Test ID
        'B': 20,  # Domain
        'C': 10,  # Priority
        'D': 18,  # Test Role
        'E': 14,  # Location
        'F': 28,  # Test Scenario
        'G': 48,  # Steps / Input
        'H': 48,  # Expected Result
        'I': 14,  # Status
        'J': 24,  # Evidence / URL
        'K': 14,  # Tester
        'L': 14,  # Test Date
        'M': 32   # Notes / Bug ID
    }
    for col_l, width in uat_col_widths.items():
        ws_uat.column_dimensions[col_l].width = width

    mob_col_widths = {
        'A': 12,  # Test ID
        'B': 18,  # Domain
        'C': 10,  # Priority
        'D': 18,  # Test Role
        'E': 16,  # Location
        'F': 28,  # Scenario
        'G': 44,  # Steps / Input
        'H': 46,  # Expected Result
        'I': 24,  # Device / OS / Browser
        'J': 14,  # Status
        'K': 22,  # Evidence / Link
        'L': 14,  # Tester
        'M': 14,  # Test Date
        'N': 30   # Catatan / Bug ID
    }
    for col_l, width in mob_col_widths.items():
        ws_mobile.column_dimensions[col_l].width = width

    # Save workbook
    wb.save(output_path)
    print(f"QA Workbook successfully created at: {output_path}")
    print(f"Total UAT Desktop Test Cases: {len(uat_cases)}")
    print(f"Total Mobile Test Cases: {len(mob_cases)}")
    print(f"Grand Total Test Cases: {len(uat_cases) + len(mob_cases)}")

if __name__ == '__main__':
    build_complete_qa_workbook()
